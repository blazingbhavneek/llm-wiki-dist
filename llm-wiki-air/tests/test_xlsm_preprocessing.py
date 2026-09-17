from __future__ import annotations

import asyncio
import tempfile
import unittest
import zipfile
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference

from graph.formats.tabular import SheetStructure, TableSpec, VbaDescription, _table_origin, decide_structure, find_regions, grid_from_html, write_tables
from graph.formats.xlsx import _append_story, _cells_for_ranges, _replace_line_references, _story_source, split_sheets
from graph.wiki.config import WikiConfig
from graph.wiki.page import check_section
from graph.workspace.parser_client import parse_document
from graph.workspace.xlsm import apply_manifest, build_manifest


class XlsmPreprocessingTests(unittest.TestCase):
    def _workbook(self, path: Path, *, macro: bool) -> None:
        workbook = Workbook()
        settings = workbook.active
        settings.title = "設定"
        settings["A1"] = 1
        aggregate = workbook.create_sheet("集計")
        aggregate.sheet_state = "hidden"
        aggregate["A1"] = "=設定!A1"
        search = workbook.create_sheet("検索結果")
        search.sheet_state = "hidden"
        search["A1"] = 2
        data = workbook.create_sheet("グラフデータ")
        data.sheet_state = "hidden"
        for row in range(1, 244):
            data.cell(row, 1, "=集計!A1+検索結果!A1")
        graph = workbook.create_sheet("グラフ")
        chart = BarChart()
        chart.add_data(Reference(data, min_col=1, min_row=1, max_row=243))
        graph.add_chart(chart, "A1")
        workbook.save(path)
        workbook.close()
        if macro:
            with zipfile.ZipFile(path, "a") as archive:
                archive.writestr("xl/vbaProject.bin", b"test")

    def test_only_macro_workbooks_receive_a_sheet_lineage_manifest(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            plain = Path(directory) / "plain.xlsm"
            macro = Path(directory) / "macro.xlsm"
            self._workbook(plain, macro=False)
            self._workbook(macro, macro=True)
            code = (
                'Sub BuildGraph()\nSet src = Worksheets("検索結果")\n'
                'Set dst = Worksheets("グラフデータ")\n'
                'dst.Cells(1, 1).Value = src.Cells(1, 1).Value\nEnd Sub\n'
                'Sub DynamicSheet()\nSet ws = Worksheets(sheetName)\nEnd Sub\n'
            )
            with patch("graph.workspace.xlsm._modules", return_value=[("Module1.bas", code)]):
                manifest = build_manifest(macro)

            self.assertIsNone(build_manifest(plain))
            rows = {row["name"]: row for row in manifest["sheets"]}
            self.assertEqual(rows["グラフ"]["emit"], "full")
            self.assertEqual(rows["グラフデータ"]["emit"], "lineage")
            self.assertEqual(rows["グラフデータ"]["role"], "hidden_intermediate")
            self.assertIn("グラフデータ", rows["グラフ"]["depends_on"])
            self.assertIn("設定", rows["グラフ"]["lineage"])
            build = next(row for row in manifest["procedures"] if row["name"] == "BuildGraph")
            self.assertEqual(build["sheets_read"], ["検索結果"])
            self.assertEqual(build["sheets_written"], ["グラフデータ"])
            self.assertIn("グラフ", build["final_outputs_affected"])
            dynamic = next(row for row in manifest["procedures"] if row["name"] == "DynamicSheet")
            self.assertTrue(dynamic["unresolved_dynamic_references"])

    @patch("graph.workspace.parser_client.requests.post")
    def test_non_xlsm_upload_has_no_manifest(self, post: Mock) -> None:
        post.return_value.status_code = 200
        post.return_value.json.return_value = {"markdown": "ok"}
        post.return_value.raise_for_status.return_value = None
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.pdf"
            source.write_bytes(b"pdf")
            parse_document(source, base_url="http://parser", settings=object())
        self.assertIsNone(post.call_args.kwargs["data"])

    @patch("graph.workspace.parser_client.requests.post")
    @patch("graph.workspace.parser_client.build_manifest")
    def test_macro_upload_sends_original_and_manifest(self, manifest: Mock, post: Mock) -> None:
        manifest.return_value = {"mode": "xlsm-vba-lineage"}
        post.return_value.status_code = 200
        post.return_value.json.return_value = {"markdown": "ok"}
        post.return_value.raise_for_status.return_value = None
        with tempfile.TemporaryDirectory() as directory:
            source = Path(directory) / "source.xlsm"
            source.write_bytes(b"original-xlsm")
            parse_document(source, base_url="http://parser", settings=object())
        self.assertEqual(
            post.call_args.kwargs["data"],
            {"manifest": '{"mode": "xlsm-vba-lineage"}'},
        )
        self.assertEqual(post.call_args.kwargs["files"]["file"][0], "source.xlsm")

    def test_sheet_context_survives_table_page_extraction(self) -> None:
        sheets = split_sheets(
            [
                "# workbook",
                "## シート: グラフ",
                "<table><tr><td>result</td></tr></table>",
                "<!-- sheet-context:start -->",
                "### Sheet lineage",
                "- All upstream sheets: `グラフデータ`",
                "<!-- sheet-context:end -->",
            ]
        )
        self.assertEqual(len(sheets), 1)
        self.assertIn("All upstream sheets: `グラフデータ`", sheets[0][1])

    def test_legacy_parser_vba_block_is_replaced_and_ordered_after_sheets(self) -> None:
        manifest = {
            "sheets": [{"name": "操作", "lineage": [], "charts": []}],
            "procedures": [{
                "id": "module1::run",
                "module": "Module1.bas",
                "name": "Run",
                "kind": "sub",
                "code": "Sub Run()\nEnd Sub",
                "buttons": [{"sheet": "操作", "cells": "B2:D4"}],
                "formula_cells": [],
                "sheets_read": ["操作"],
                "sheets_written": [],
                "final_outputs_affected": [],
                "unresolved_dynamic_references": [],
            }],
        }
        normalized = apply_manifest(
            "# book\n\n## シート: VBA\n\n```vb\nall code\n```\n\n"
            "## シート: 操作\n\n<table><tr><td>x</td></tr></table>\n",
            manifest,
        )
        sheets = split_sheets(normalized.splitlines())

        self.assertEqual([name for name, _, _ in sheets], ["操作", "マクロ-Module1-Run"])
        self.assertIn("[Run](002-マクロ-Module1-Run.md)", sheets[0][1])
        self.assertIn("```vb\nSub Run()\nEnd Sub\n```", sheets[1][1])

    def test_story_provenance_uses_cells_instead_of_parser_lines(self) -> None:
        source, spans = _story_source([
            ("操作-part1", "_元範囲: A1:J100_\n\n<table></table>", (20, 40)),
            ("操作-part2", "_元範囲: A101:J180_\n\n<table></table>", (41, 60)),
        ])
        self.assertEqual(len(source.split("\n")) - 1, 2)
        self.assertIn("参照セル範囲: 操作!A1:J100", source)
        self.assertEqual(_cells_for_ranges([[1, 1]], spans), ["操作!A1:J100"])
        self.assertNotIn("原文 1-1行", _replace_line_references("原文 1-1行", spans))

    def test_regular_excel_table_keeps_its_cell_coordinates(self) -> None:
        body = "<table><tr><td>見出し</td><td>値</td></tr><tr><td>項目</td><td>7</td></tr></table>\n\n_元範囲: B3:C4_"
        grid = grid_from_html(body, origin=_table_origin(body))
        self.assertEqual(grid.get(3, 2), "見出し")
        self.assertEqual(grid.get(4, 3), "7")


class _Model:
    async def structured(self, schema, _messages):
        if schema is VbaDescription:
            return VbaDescription(title="集計実行", summary="集計を実行します。")
        return SheetStructure(
            summary="summary",
            tables=[TableSpec(region=1, title="Table", data_rows=[1, 2], data_cols=[1, 1])],
        )

    async def text(self, _messages):
        return "analysis"


class XlsmPageOrderTests(unittest.IsolatedAsyncioTestCase):
    def test_excel_story_does_not_require_every_cell_address(self) -> None:
        errors = check_section(
            "要点を説明する本文",
            lines=["AA29 AB30 AC31"],
            source_text="AA29 AB30 AC31",
            block_ranges=[],
            placeholders=[],
            facts=[],
            check_identifiers=False,
        )

        self.assertEqual(errors, [])

    async def test_sheet_structure_asks_only_for_orientation(self) -> None:
        class ChoiceModel:
            calls = 0

            async def text(self, _messages):
                self.calls += 1
                return "列"

            async def structured(self, *_args, **_kwargs):
                raise AssertionError("structured generation must not be used")

        model = ChoiceModel()
        grid = grid_from_html(
            "<table><tr><th>項目</th><th>一月</th></tr><tr><td>売上</td><td>10</td></tr></table>"
        )
        structure = await decide_structure(
            "集計",
            grid,
            find_regions(grid),
            model=model,
            config=SimpleNamespace(
                tabular_preview_rows=12,
                tabular_preview_cols=12,
                source_kind="xlsx",
            ),
        )

        self.assertEqual(model.calls, 1)
        self.assertEqual(structure.tables[0].orientation, "columns")
        self.assertEqual(structure.tables[0].data_cols, [2, 2])

    async def test_source_descriptions_use_configured_concurrency(self) -> None:
        class ConcurrentModel:
            active = 0
            maximum = 0

            async def text(self, _messages):
                self.active += 1
                self.maximum = max(self.maximum, self.active)
                await asyncio.sleep(0.01)
                self.active -= 1
                return "行"

        model = ConcurrentModel()
        table = "<table><tr><td>A</td></tr><tr><td>1</td></tr></table>"
        with tempfile.TemporaryDirectory() as directory:
            await write_tables(
                sheets=[(f"Sheet{i}", table, (i, i)) for i in range(4)],
                run_dir=Path(directory),
                model=model,
                config=SimpleNamespace(
                    tabular_preview_rows=12,
                    tabular_preview_cols=12,
                    tabular_slice_records=40,
                    output_language="Japanese",
                    source_kind="xlsx",
                    rewrite_concurrency=2,
                ),
                generate_analyses=False,
            )
            names = sorted(path.name for path in (Path(directory) / "docs").glob("*.md"))

        self.assertEqual(model.maximum, 2)
        self.assertEqual(names, [f"{i:03d}-シート-Sheet{i - 1}.md" for i in range(1, 5)])

    async def test_workbook_story_is_appended_with_cell_provenance(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            output = root / "output"
            (output / "docs").mkdir(parents=True)
            source = root / "book.md"
            source.write_text("source", encoding="utf-8")
            files = [{
                "filename": "001-シート-操作.md",
                "title": "操作",
                "kind": "table",
                "source_cells": ["操作!A1:J10"],
                "summary": "原資料",
            }]

            async def fake_pipeline(_source, *, config, **_kwargs):
                run = Path(config.run_dir)
                (run / "state").mkdir(parents=True)
                (run / "wiki").mkdir(parents=True)
                (run / "state" / "plan.json").write_text(json.dumps({"pages": [{
                    "filename": "001-操作の流れ.md",
                    "title": "操作の流れ",
                    "summary": "全体説明",
                    "owner_ranges": [[1, 1]],
                }]}, ensure_ascii=False), encoding="utf-8")
                (run / "wiki" / "001-操作の流れ.md").write_text("# 操作の流れ\n\n原文 1-1行", encoding="utf-8")
                return run

            with patch("graph.wiki.pipeline.run_pipeline", side_effect=fake_pipeline):
                await _append_story(
                    source,
                    [("操作", "_元範囲: A1:J10_\n\n<table></table>", (1, 3))],
                    files,
                    run_dir=output,
                    model=object(),
                    config=WikiConfig(run_dir=str(root / "state"), resume=False),
                )

            manifest = json.loads((output / "_planning" / "manifest.json").read_text(encoding="utf-8"))
            story = (output / "docs" / "002-解説-操作の流れ.md").read_text(encoding="utf-8")
            self.assertEqual(manifest["files"][1]["source_cells"], ["操作!A1:J10"])
            self.assertNotIn("source_ranges", json.dumps(manifest))
            self.assertIn("参照セル `操作!A1:J10`", story)

    async def test_sheet_then_vba_then_series_numbering(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            await write_tables(
                sheets=[
                    ("Sheet1", "<table><tr><th>Row</th><th>A</th></tr><tr><td>1</td><td>x</td></tr><tr><td>2</td><td>y</td></tr></table>\n\n[Run](002-マクロ-Module-Run.md)", (1, 3)),
                    ("vba-Module-Run", "<!-- vba-id: module%3A%3Arun -->\n```vb\nSub Run()\nEnd Sub\n```", (4, 7)),
                ],
                run_dir=Path(directory),
                model=_Model(),
                config=SimpleNamespace(
                    tabular_preview_rows=12,
                    tabular_preview_cols=12,
                    tabular_slice_records=40,
                    output_language="Japanese",
                    source_kind="xlsx",
                ),
            )
            names = sorted(path.name for path in (Path(directory) / "docs").glob("*.md"))
            sheet = (Path(directory) / "docs" / "001-シート-Sheet1.md").read_text(encoding="utf-8")

        self.assertEqual(names[:2], ["001-シート-Sheet1.md", "002-マクロ-集計実行.md"])
        self.assertEqual(names[2], "003-解説-Sheet1-領域-1-分析.md")
        self.assertIn("[Run](002-マクロ-集計実行.md)", sheet)


if __name__ == "__main__":
    unittest.main()
