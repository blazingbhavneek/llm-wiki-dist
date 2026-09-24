"""Acceptance tests for the Git-backed mount diff pipeline.

Fast fixture checks run during ordinary unittest discovery.  The acceptance
tests are intentionally opt-in because they call the configured parser, LLM,
embedding service, and GROWI instance.

Run all acceptance cases:

    RUN_DIFF_INTEGRATION=1 .venv/bin/python -m unittest -v \
        tests.test_mount_diff_pipeline.MountDiffPipelineAcceptanceTest

Run one case by appending its complete method name to the command above.
Every acceptance case creates a temporary local project and a unique
``/diff-test-*`` GROWI namespace, then removes its remote pages in cleanup.
"""

from __future__ import annotations

import asyncio
import contextlib
import hashlib
import json
import math
import os
import re
import shutil
import sqlite3
import subprocess
import tempfile
import unittest
import uuid
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from docx import Document

from graph import config
from graph.config import Settings
from graph.growi.client import GrowiPage, GrowiPublisher
from graph.workspace.project import Project, open_project, raw_name_for
from publisher import pipeline, queue
from publisher.history import candidate, ensure_repository, last_good, read_blob, stage_blob
from publisher.ledger import Ledger, load_ledger, save_ledger


ROOT = Path(__file__).resolve().parents[1]
CONFIG = ROOT / "configs" / "diff_test_local.ini"
FIXTURE = ROOT / "tests" / "samples" / "digital-agency-pdl.docx"
RUN_LIVE = os.environ.get("RUN_DIFF_INTEGRATION") == "1"
LAST_GOOD_REF = "refs/llm-wiki/last-good"


def _nonempty_paragraphs(path: Path):
    document = Document(path)
    return document, [paragraph for paragraph in document.paragraphs if paragraph.text.strip()]


def make_minor_docx_change(path: Path, marker: str) -> None:
    """Change one paragraph without changing the document structure."""
    document, paragraphs = _nonempty_paragraphs(path)
    if not paragraphs:
        raise AssertionError(f"DOCX fixture has no editable paragraphs: {path}")
    paragraphs[len(paragraphs) // 2].text += f" [{marker}]"
    document.save(path)


def make_major_docx_change(path: Path, marker: str, fraction: float = 0.40) -> None:
    """Rewrite enough distributed paragraphs to exceed either size threshold."""
    document, paragraphs = _nonempty_paragraphs(path)
    if not paragraphs:
        raise AssertionError(f"DOCX fixture has no editable paragraphs: {path}")
    count = max(1, math.ceil(len(paragraphs) * fraction))
    # Spread changes through the document so locality cannot classify this as
    # one small edit even if parser line boundaries differ from DOCX paragraphs.
    indexes = [round(i * (len(paragraphs) - 1) / max(count - 1, 1)) for i in range(count)]
    for sequence, index in enumerate(dict.fromkeys(indexes), start=1):
        paragraphs[index].text = f"{marker} rewritten paragraph {sequence}"
    document.save(path)


class DiffDocxFixtureTest(unittest.TestCase):
    def test_diff_project_config_and_docx_fixture_exist(self) -> None:
        self.assertTrue(CONFIG.is_file(), CONFIG)
        self.assertTrue(FIXTURE.is_file(), FIXTURE)
        self.assertGreater(FIXTURE.stat().st_size, 0)
        config.PROJECT_ROOT = ROOT
        settings = Settings.from_env(str(CONFIG))
        self.assertEqual(settings.mount_path, "/tmp/llm-wiki-diff-test/mount")
        self.assertEqual(settings.target_name, "diff_test_local")
        self.assertEqual(settings.chat_base_url, "http://127.0.0.1:8080/v1")
        self.assertEqual(settings.parser_base_url, "http://127.0.0.1:8000/agent/doc-parser/")

    def test_docx_mutators_create_small_and_large_different_files(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            minor = Path(tmp) / "minor.docx"
            major = Path(tmp) / "major.docx"
            shutil.copy2(FIXTURE, minor)
            shutil.copy2(FIXTURE, major)
            make_minor_docx_change(minor, "minor-fixture-check")
            make_major_docx_change(major, "major-fixture-check")
            self.assertNotEqual(FIXTURE.read_bytes(), minor.read_bytes())
            self.assertNotEqual(FIXTURE.read_bytes(), major.read_bytes())
            self.assertNotEqual(minor.read_bytes(), major.read_bytes())


class DiffPipelineSafetyTest(unittest.TestCase):
    def test_one_source_edit_can_update_repeated_generated_facts(self) -> None:
        from graph.wiki.wire import IncrementalPageEditResult
        from graph.workspace.writer import _apply_model_patches

        result = IncrementalPageEditResult.model_validate({"patches": [
            {"edit_ids": [1], "before": "概要はＴ１分。", "after": "概要はＴ２分。"},
            {"edit_ids": [1], "before": "詳細もＴ１分。", "after": "詳細もＴ２分。"},
        ]})

        updated = _apply_model_patches("概要はＴ１分。\n\n詳細もＴ１分。\n", result, {1})

        self.assertEqual(updated, "概要はＴ２分。\n\n詳細もＴ２分。\n")

    def test_many_distributed_small_edits_stay_small(self) -> None:
        old = [f"line {index}" for index in range(2000)]
        new = list(old)
        for index in range(0, 2000, 20):
            new[index] += " updated"

        with patch.object(queue, "_docx_text", side_effect=[old, new]):
            result = queue._classification(b"old", b"new", ".docx")

        self.assertEqual(result["kind"], "small")
        self.assertEqual(result["hunks"], 100)
        self.assertEqual(result["reason"], "below-threshold")

    def test_sync_command_rescans_until_the_queue_is_drained(self) -> None:
        import main

        settings = SimpleNamespace(data_root="/tmp/data", target_name="test", mount_path="/tmp/mount")
        args = SimpleNamespace(items=["test.docx"], force=True, verbose=False)
        first = {"done": [{"path": "test.docx", "status": "added"}], "failures": []}
        second = {"done": [{"path": "test.docx", "status": "changed"}], "failures": []}
        project = object()
        with (
            patch.object(main, "_settings", return_value=settings),
            patch.object(main, "open_project", return_value=project),
            patch("publisher.queue.worker_lock", return_value=contextlib.nullcontext()),
            patch("publisher.queue.retry_failed") as retry,
            patch("publisher.queue.scan") as scan,
            patch("publisher.queue.work_once", side_effect=[first, second, None]) as work,
        ):
            result = main.cmd_sync(args)

        self.assertEqual(result, 0)
        retry.assert_called_once_with(project)
        self.assertEqual(work.call_count, 3)
        self.assertEqual(scan.call_count, 3)
        self.assertTrue(scan.call_args_list[0].kwargs["force"])
        self.assertFalse(scan.call_args_list[1].kwargs["force"])
        self.assertTrue(all(call.kwargs["verify_content"] for call in scan.call_args_list))

    def test_small_paragraph_deletion_uses_model_and_scopes_linker(self) -> None:
        from graph.workspace import writer

        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            mount = root / "mount"
            mount.mkdir()
            project = Project(root / "project", mount).ensure()
            rel = "test_docx.md"
            prefix = "# Title\n\n" + "\n".join(f"line {index}" for index in range(300))
            old_source = prefix + "\n\n本機能が主に動作する計算機・装置を示す。\n\nend\n"
            new_source = prefix + "\n\nend\n"
            project.raw_file(rel).write_text(new_source, encoding="utf-8")
            state = project.state_dir(rel)
            (state / "source").mkdir(parents=True)
            (state / "source" / "original.md").write_text(old_source, encoding="utf-8")
            (state / "state" / "pages").mkdir(parents=True)
            (state / "state" / "plan.json").write_text(
                json.dumps({
                    "source_line_count": len(old_source.splitlines()),
                    "pages": [{
                        "number": 1,
                        "filename": "001.md",
                        "title": "Title",
                        "chapter": "",
                        "path": [],
                        "summary": "",
                        "owner_ranges": [[1, len(old_source.splitlines())]],
                        "reference_ranges": [],
                    }],
                }),
                encoding="utf-8",
            )
            (state / "state" / "pages" / "001.json").write_text("{}", encoding="utf-8")
            (state / "wiki").mkdir()
            stale = (
                "# Title\n\nkeep\n\n"
                "本機能が主に動作する計算機および装置の構成について述べる。\n\n"
                "## End\n\nend\n"
            )
            (state / "wiki" / "001.md").write_text(stale, encoding="utf-8")
            target = project.wiki_dir(rel)
            (target / "_planning").mkdir(parents=True)
            (target / "_planning" / "linker.json").write_text(
                json.dumps({"schema_version": 2, "status": "complete", "mode": "legacy"}),
                encoding="utf-8",
            )

            class FakeModel:
                def __init__(self):
                    self.prompts: list[str] = []

                async def structured(self, schema, messages, **_kwargs):
                    self.prompts.append("\n".join(map(str, messages)))
                    return schema.model_validate({
                        "patches": [{
                            "edit_ids": [1],
                            "before": "本機能が主に動作する計算機および装置の構成について述べる。",
                            "after": "",
                        }],
                    })

            settings = SimpleNamespace(
                ingest_mode="wiki",
                wiki_linker_enabled=True,
                wiki_linker_mode="legacy",
            )
            model = FakeModel()
            with patch.object(writer, "build_wiki_output", side_effect=AssertionError("full wiki runner invoked")):
                result = writer.write_wiki_pages(
                    project, rel, mode="wiki", settings=settings,
                    llm=model, embedder=None,
                )

            self.assertEqual(result.tier, 1)
            self.assertEqual(result.rebuild, "incremental")
            page = (project.wiki_dir(rel) / "001.md").read_text(encoding="utf-8")
            self.assertNotIn("本機能が主に動作する計算機および装置", page)
            self.assertIn("## End", page)
            self.assertEqual((state / "wiki" / "001.md").read_text(encoding="utf-8"), page)
            self.assertIn("DELETE", model.prompts[0])
            self.assertIn("本機能が主に動作する計算機および装置", model.prompts[0])
            marker = json.loads((project.wiki_dir(rel) / "_planning" / "linker.json").read_text(encoding="utf-8"))
            self.assertEqual(marker["status"], "pending")
            self.assertTrue(marker["resume"])
            with patch.object(writer, "build_wiki_output", side_effect=AssertionError("full wiki runner invoked")):
                unchanged = writer.write_wiki_pages(
                    project, rel, mode="wiki", settings=settings,
                    llm=model, embedder=None,
                )
            self.assertEqual(unchanged.changed_pages, [])
            self.assertEqual(len(model.prompts), 1)

    def test_rewrite_prompts_prohibit_reviving_deleted_source_content(self) -> None:
        from graph.wiki.prompts import incremental_page_edit_prompt, intro_prompt, section_write_prompt

        section = section_write_prompt(
            page_title="試験", page_summary="", index=1, count=1,
            source_start=1, source_end=1, numbered_section="1: 現行本文",
            facts_text="なし", image_context="なし", output_language="日本語",
        ).render()
        intro = intro_prompt(
            page_title="試験", page_summary="現行要約", body="現行本文",
            output_language="日本語",
        ).render()
        self.assertIn("削除された段落", section)
        self.assertIn("削除された内容を復活させない", intro)
        incremental = incremental_page_edit_prompt(
            page_title="試験", current_page="# 試験\n\n旧本文", current_source="1: 新本文",
            edits="## EDIT 1: DELETE\n旧本文", image_context="なし", output_language="日本語",
        ).render()
        self.assertIn("ADD、UPDATE、DELETEを一件も漏らさず", incremental)
        self.assertIn("言い換えられていても対応する内容を完全に削除", incremental)

    def test_incremental_linker_keeps_valid_peer_untouched_and_skips_search(self) -> None:
        from graph.linker.catalog import Catalog
        from graph.linker.chunks import make_chunks, to_json
        from graph.linker.prompts import CHUNK_META_VERSION
        from graph.linker.service import link_document
        from graph.linker.wire import EdgeSuggestions
        from graph.wiki.storage import write_json_atomic

        with tempfile.TemporaryDirectory() as tmp:
            project = Project(Path(tmp) / "project", Path(tmp) / "mount").ensure()
            rel = "doc.md"
            document = "doc"
            old_text = "# Doc\n\nold relation\n"
            new_text = "# Doc\n\nnew relation\n"
            peer_text = "# Peer\n\npeer\n"
            old_chunk = make_chunks(document, "general", "001.md", old_text, id_seed=rel)[0]
            peer_chunk = make_chunks("peer", "general", "001.md", peer_text, id_seed="peer.md")[0]
            catalog = Catalog.open(project.linker_database, mode="legacy")
            try:
                catalog.reconcile(document, [old_chunk], team="general", raw_rel=rel)
                catalog.reconcile("peer", [peer_chunk], team="general", raw_rel="peer.md")
                catalog.insert_edge({
                    "chunk_a": old_chunk.chunk_id,
                    "chunk_b": peer_chunk.chunk_id,
                    "label": "uses",
                    "summary": "existing",
                    "source": "legacy_rrf",
                })
            finally:
                catalog.close()

            planning = project.wiki_dir(rel) / "_planning"
            (planning / "pages").mkdir(parents=True)
            (planning / "pages" / "001.md").write_text(new_text, encoding="utf-8")
            write_json_atomic(planning / "source.json", {"id_seed": rel})
            write_json_atomic(planning / "linker.json", {"status": "pending", "resume": True, "mode": "legacy"})
            cached = to_json(document, "general", [old_chunk], id_seed=rel)
            cached["meta_version"] = CHUNK_META_VERSION
            write_json_atomic(planning / "chunks.json", cached)

            class FakeModel:
                async def structured(self, schema, _messages, **_kwargs):
                    if schema is EdgeSuggestions:
                        return EdgeSuggestions(edges=[{
                            "target_node_id": peer_chunk.chunk_id,
                            "label": "uses",
                            "summary": "existing",
                        }])
                    raise AssertionError(schema)

            settings = SimpleNamespace(
                wiki_linker_mode="legacy", wiki_output_language="日本語",
                wiki_linker_concurrency=1, concurrency=1,
            )
            with patch("graph.linker.legacy.candidates", side_effect=AssertionError("global search ran")):
                result = asyncio.run(link_document(
                    project, rel, model=FakeModel(), embedder=None, settings=settings,
                    render=False, changed_pages={"doc/001.md"},
                ))

            self.assertEqual(result.meta_calls, 0)
            self.assertEqual(result.edge_calls, 1)
            self.assertEqual(set(result.affected_pages or []), {"doc/001.md"})

            catalog = Catalog.open(project.linker_database, mode="legacy")
            try:
                edge = catalog.edges_for_page("doc/001.md")[0]
                self.assertEqual((edge["label"], edge["summary"]), ("uses", "existing"))
            finally:
                catalog.close()

            (planning / "pages" / "001.md").write_text("# Doc\n\nno relation\n", encoding="utf-8")
            write_json_atomic(planning / "linker.json", {"status": "pending", "resume": True, "mode": "legacy"})

            class RejectModel(FakeModel):
                async def structured(self, schema, messages, **kwargs):
                    if schema is EdgeSuggestions:
                        return EdgeSuggestions(edges=[])
                    return await super().structured(schema, messages, **kwargs)

            with patch("graph.linker.legacy.candidates", side_effect=AssertionError("global search ran")):
                rejected = asyncio.run(link_document(
                    project, rel, model=RejectModel(), embedder=None, settings=settings,
                    render=False, changed_pages={"doc/001.md"},
                ))

            self.assertEqual(set(rejected.affected_pages or []), {"doc/001.md", "peer/001.md"})
            self.assertGreaterEqual(rejected.edges_removed, 1)

    def test_neo_incremental_linker_preserves_entity_direction_and_checks_only_visible_edges(self) -> None:
        from graph.linker.catalog import Catalog
        from graph.linker.chunks import make_chunks, to_json
        from graph.linker.prompts import CHUNK_META_VERSION
        from graph.linker.service import link_document
        from graph.linker.wire import ChunkMeta, NeoEdgeSuggestions
        from graph.wiki.storage import write_json_atomic

        with tempfile.TemporaryDirectory() as tmp:
            project = Project(Path(tmp) / "project", Path(tmp) / "mount").ensure()
            rel = "doc.md"
            old_chunk = make_chunks("doc", "general", "001.md", "# Doc\n\nold\n", id_seed=rel)[0]
            old_chunk.meta = ChunkMeta(summary="doc", entities=[{"name": "Shared", "kind": "concept", "role": "defines"}])

            def peer(document: str, text: str, *, uses: bool = False):
                item = make_chunks(document, "general", "001.md", f"# {document}\n\n{text}\n", id_seed=f"{document}.md")[0]
                item.meta = ChunkMeta(
                    summary=document,
                    entities=[{"name": "Shared", "kind": "concept", "role": "uses"}] if uses else [],
                )
                return item

            user = peer("user", "Shared", uses=True)
            visible = peer("visible", "visible relation")
            hidden = peer("hidden", "hidden relation")
            catalog = Catalog.open(project.linker_database, mode="neo")
            try:
                for document, raw_rel, item in (
                    ("doc", rel, old_chunk),
                    ("user", "user.md", user),
                    ("visible", "visible.md", visible),
                    ("hidden", "hidden.md", hidden),
                ):
                    catalog.reconcile(document, [item], team="general", raw_rel=raw_rel)
                catalog.insert_edge({
                    "chunk_a": old_chunk.chunk_id, "chunk_b": user.chunk_id,
                    "label": "uses", "summary": "entity", "source": "define", "via": ["Shared"],
                })
                catalog.insert_edge({
                    "chunk_a": visible.chunk_id, "chunk_b": old_chunk.chunk_id,
                    "label": "prerequisite", "summary": "visible", "source": "hop1",
                })
                catalog.insert_edge({
                    "chunk_a": hidden.chunk_id, "chunk_b": old_chunk.chunk_id,
                    "label": "prerequisite", "summary": "hidden", "source": "hop2",
                })
                baseline = {str(row["summary"]): dict(row) for row in catalog.conn.execute("SELECT * FROM edges")}
            finally:
                catalog.close()

            planning = project.wiki_dir(rel) / "_planning"
            (planning / "pages").mkdir(parents=True)
            (planning / "pages" / "001.md").write_text("# Doc\n\nchanged Shared\n", encoding="utf-8")
            write_json_atomic(planning / "source.json", {"id_seed": rel})
            write_json_atomic(planning / "linker.json", {"status": "pending", "resume": True, "mode": "neo"})
            cached = to_json("doc", "general", [old_chunk], id_seed=rel)
            cached["meta_version"] = CHUNK_META_VERSION
            write_json_atomic(planning / "chunks.json", cached)
            visible_planning = project.wiki / "visible" / "_planning"
            visible_planning.mkdir(parents=True)
            write_json_atomic(visible_planning / "navigation.json", {
                "pages": {"001.md": {"references": [{"edge_id": baseline["visible"]["edge_id"]}]}}
            })

            class Model:
                def __init__(self, accept: bool = True):
                    self.accept = accept
                    self.edge_targets: list[str] = []

                async def structured(self, schema, messages, **_kwargs):
                    if schema is ChunkMeta:
                        raise AssertionError("incremental metadata regenerated")
                    if schema is NeoEdgeSuggestions:
                        payload = json.loads(messages[-1].content)
                        self.edge_targets.append(payload["target"]["chunk_id"])
                        candidates = payload["candidates"] if self.accept else []
                        return NeoEdgeSuggestions(edges=[{
                            "target_chunk_id": item["chunk_id"],
                            "label": "prerequisite", "summary": "still valid",
                        } for item in candidates])
                    raise AssertionError(schema)

            settings = SimpleNamespace(
                wiki_linker_mode="neo", wiki_output_language="日本語",
                wiki_linker_concurrency=1, concurrency=1,
            )
            model = Model()
            with patch("graph.linker.neo.candidates", side_effect=AssertionError("global search ran")):
                result = asyncio.run(link_document(
                    project, rel, model=model, embedder=None, settings=settings,
                    render=False, changed_pages={"doc/001.md"},
                ))

            self.assertEqual(result.edge_calls, 1)
            self.assertEqual(result.meta_calls, 0)
            self.assertEqual(model.edge_targets, [visible.chunk_id])
            self.assertEqual(set(result.affected_pages or []), {"doc/001.md"})
            catalog = Catalog.open(project.linker_database, mode="neo")
            try:
                current = {str(row["summary"]): dict(row) for row in catalog.conn.execute("SELECT * FROM edges")}
            finally:
                catalog.close()
            self.assertEqual(set(current), set(baseline))
            for summary in baseline:
                self.assertEqual(
                    (current[summary]["edge_id"], current[summary]["chunk_a"], current[summary]["chunk_b"]),
                    (baseline[summary]["edge_id"], baseline[summary]["chunk_a"], baseline[summary]["chunk_b"]),
                )

            (planning / "pages" / "001.md").write_text("# Doc\n\nchanged again\n", encoding="utf-8")
            write_json_atomic(planning / "linker.json", {"status": "pending", "resume": True, "mode": "neo"})
            rejected = asyncio.run(link_document(
                project, rel, model=Model(False), embedder=None, settings=settings,
                render=False, changed_pages={"doc/001.md"},
            ))
            self.assertEqual(set(rejected.affected_pages or []), {"doc/001.md", "user/001.md", "visible/001.md"})
            catalog = Catalog.open(project.linker_database, mode="neo")
            try:
                remaining = {str(row["summary"]) for row in catalog.conn.execute("SELECT * FROM edges")}
            finally:
                catalog.close()
            self.assertEqual(remaining, {"hidden"})

    def test_regenerated_page_is_described_and_searched_only_for_its_chunks(self) -> None:
        from graph.linker.catalog import Catalog
        from graph.linker.chunks import make_chunks, to_json
        from graph.linker.legacy import Candidate
        from graph.linker.prompts import CHUNK_META_VERSION
        from graph.linker.service import link_document
        from graph.linker.wire import ChunkMeta, EdgeSuggestions
        from graph.wiki.storage import write_json_atomic

        with tempfile.TemporaryDirectory() as tmp:
            project = Project(Path(tmp) / "project", Path(tmp) / "mount").ensure()
            rel = "doc.md"
            old = make_chunks("doc", "general", "001.md", "# Doc\n\nold relation\n", id_seed=rel)[0]
            peer = make_chunks("peer", "general", "001.md", "# Peer\n\npeer relation\n", id_seed="peer.md")[0]
            catalog = Catalog.open(project.linker_database, mode="legacy")
            try:
                catalog.reconcile("doc", [old], team="general", raw_rel=rel)
                catalog.reconcile("peer", [peer], team="general", raw_rel="peer.md")
                catalog.insert_edge({
                    "chunk_a": old.chunk_id, "chunk_b": peer.chunk_id,
                    "label": "uses", "summary": "old relation", "source": "legacy_rrf",
                })
            finally:
                catalog.close()

            new_text = "# Doc\n\nnewly regenerated relation\n"
            new = make_chunks("doc", "general", "001.md", new_text, id_seed=rel)[0]
            planning = project.wiki_dir(rel) / "_planning"
            (planning / "pages").mkdir(parents=True)
            (planning / "pages" / "001.md").write_text(new_text, encoding="utf-8")
            write_json_atomic(planning / "source.json", {"id_seed": rel})
            write_json_atomic(planning / "linker.json", {"status": "pending", "resume": True, "mode": "legacy"})
            cached = to_json("doc", "general", [old], id_seed=rel)
            cached["meta_version"] = CHUNK_META_VERSION
            write_json_atomic(planning / "chunks.json", cached)
            searched: list[str] = []

            class Model:
                async def structured(self, schema, _messages, **_kwargs):
                    if schema is ChunkMeta:
                        return ChunkMeta(summary="new", keywords=["k"])
                    if schema is EdgeSuggestions:
                        return EdgeSuggestions(edges=[{
                            "target_node_id": peer.chunk_id,
                            "label": "uses", "summary": "new relation",
                        }])
                    raise AssertionError(schema)

            def candidates(_catalog, item, **_kwargs):
                searched.append(item.chunk_id)
                return [Candidate(peer.chunk_id, "legacy_rrf")]

            settings = SimpleNamespace(
                wiki_linker_mode="legacy", wiki_output_language="日本語",
                wiki_linker_concurrency=1, concurrency=1,
            )
            with patch("graph.linker.legacy.candidates", side_effect=candidates):
                result = asyncio.run(link_document(
                    project, rel, model=Model(), embedder=None, settings=settings,
                    render=False, changed_pages={"doc/001.md"}, regenerated_pages={"doc/001.md"},
                ))

            self.assertEqual(result.meta_calls, 1)
            self.assertEqual(searched, [new.chunk_id])
            catalog = Catalog.open(project.linker_database, mode="legacy")
            try:
                self.assertEqual(catalog.chunk(new.chunk_id)["summary"], "new")
            finally:
                catalog.close()

    def test_incremental_scope_with_many_old_edges_checks_only_visible_ones(self) -> None:
        from graph.linker.catalog import Catalog
        from graph.linker.chunks import make_chunks, to_json
        from graph.linker.prompts import CHUNK_META_VERSION
        from graph.linker.service import link_document
        from graph.linker.wire import ChunkMeta, NeoEdgeSuggestions
        from graph.wiki.storage import write_json_atomic

        with tempfile.TemporaryDirectory() as tmp:
            project = Project(Path(tmp) / "project", Path(tmp) / "mount").ensure()
            rel = "doc.md"
            old = make_chunks("doc", "general", "001.md", "# Doc\n\nold Shared relation\n", id_seed=rel)[0]
            old.meta = ChunkMeta(summary="doc", entities=[{"name": "Shared", "kind": "concept", "role": "defines"}])
            catalog = Catalog.open(project.linker_database, mode="neo")
            try:
                catalog.reconcile("doc", [old], team="general", raw_rel=rel)
                for index in range(30):
                    document = f"entity{index:02d}"
                    peer = make_chunks(document, "general", "001.md", f"# {document}\n\nShared\n", id_seed=f"{document}.md")[0]
                    peer.meta = ChunkMeta(summary=document, entities=[{"name": "Shared", "kind": "concept", "role": "uses"}])
                    catalog.reconcile(document, [peer], team="general", raw_rel=f"{document}.md")
                    catalog.insert_edge({
                        "chunk_a": old.chunk_id, "chunk_b": peer.chunk_id,
                        "label": "uses", "summary": f"entity-{index}", "source": "define", "via": ["Shared"],
                    })

                behaviour_peers = []
                for index in range(3):
                    document = f"behaviour{index}"
                    peer = make_chunks(document, "general", "001.md", f"# {document}\n\nbehaviour relation\n", id_seed=f"{document}.md")[0]
                    catalog.reconcile(document, [peer], team="general", raw_rel=f"{document}.md")
                    catalog.insert_edge({
                        "chunk_a": old.chunk_id, "chunk_b": peer.chunk_id,
                        "label": "prerequisite", "summary": f"behaviour-{index}", "source": f"hop{index + 1}",
                    })
                    behaviour_peers.append((document, peer))
                visible_id = next(
                    str(row["edge_id"]) for row in catalog.conn.execute("SELECT * FROM edges WHERE summary='behaviour-0'")
                )
            finally:
                catalog.close()

            planning = project.wiki_dir(rel) / "_planning"
            (planning / "pages").mkdir(parents=True)
            changed_text = "# Doc\n\nupdated Shared relation\n"
            (planning / "pages" / "001.md").write_text(changed_text, encoding="utf-8")
            write_json_atomic(planning / "source.json", {"id_seed": rel})
            write_json_atomic(planning / "linker.json", {"status": "pending", "resume": True, "mode": "neo"})
            cached = to_json("doc", "general", [old], id_seed=rel)
            cached["meta_version"] = CHUNK_META_VERSION
            write_json_atomic(planning / "chunks.json", cached)
            visible_doc, _ = behaviour_peers[0]
            visible_planning = project.wiki / visible_doc / "_planning"
            visible_planning.mkdir(parents=True)
            write_json_atomic(visible_planning / "navigation.json", {
                "pages": {"001.md": {"references": [{"edge_id": visible_id}]}}
            })

            class Model:
                async def structured(self, schema, _messages, **_kwargs):
                    if schema is NeoEdgeSuggestions:
                        return NeoEdgeSuggestions(edges=[])
                    if schema is ChunkMeta:
                        raise AssertionError("valid cached entity metadata should be reused")
                    raise AssertionError(schema)

            settings = SimpleNamespace(
                wiki_linker_mode="neo", wiki_output_language="日本語",
                wiki_linker_concurrency=1, concurrency=1,
            )
            result = asyncio.run(link_document(
                project, rel, model=Model(), embedder=None, settings=settings,
                render=False, changed_pages={"doc/001.md"},
            ))

            self.assertEqual(result.edge_calls, 1)
            self.assertEqual(
                set(result.affected_pages or []),
                {"doc/001.md", f"{visible_doc}/001.md"},
            )

    def test_retitled_page_rerenders_linked_peers_without_model_calls(self) -> None:
        from graph.linker.catalog import Catalog
        from graph.linker.chunks import make_chunks, to_json
        from graph.linker.prompts import CHUNK_META_VERSION, REFERENCE_PLAN_VERSION
        from graph.linker.service import link_document
        from graph.linker.wire import EdgeSuggestions, PageReferencePlan
        from graph.wiki.storage import write_json_atomic

        with tempfile.TemporaryDirectory() as tmp:
            project = Project(Path(tmp) / "project", Path(tmp) / "mount").ensure()
            rel = "doc.md"
            old_text = "# Old\n\nold body\n"
            new_text = "# New\n\nupdated body\n"
            old = make_chunks("doc", "general", "001.md", old_text, id_seed=rel)[0]
            peer = make_chunks("peer", "general", "001.md", "# Peer\n\npeer body\n", id_seed="peer.md")[0]
            catalog = Catalog.open(project.linker_database, mode="legacy")
            try:
                catalog.reconcile("doc", [old], team="general", raw_rel=rel)
                catalog.reconcile("peer", [peer], team="general", raw_rel="peer.md")
                catalog.insert_edge({
                    "chunk_a": old.chunk_id, "chunk_b": peer.chunk_id,
                    "label": "uses", "summary": "linked", "source": "legacy_rrf",
                })
                edge_id = str(catalog.conn.execute("SELECT edge_id FROM edges").fetchone()[0])
            finally:
                catalog.close()

            planning = project.wiki_dir(rel) / "_planning"
            (planning / "pages").mkdir(parents=True)
            (planning / "pages" / "001.md").write_text(new_text, encoding="utf-8")
            write_json_atomic(planning / "source.json", {"id_seed": rel})
            write_json_atomic(planning / "linker.json", {"status": "pending", "resume": True, "mode": "legacy"})
            cached = to_json("doc", "general", [old], id_seed=rel)
            cached["meta_version"] = CHUNK_META_VERSION
            write_json_atomic(planning / "chunks.json", cached)
            peer_planning = project.wiki / "peer" / "_planning"
            (peer_planning / "pages").mkdir(parents=True)
            (peer_planning / "pages" / "001.md").write_text("# Peer\n\npeer body\n", encoding="utf-8")
            state = {"version": REFERENCE_PLAN_VERSION, "candidate_ids": [edge_id], "references": [{
                "edge_id": edge_id, "placement": "footer", "anchor": "", "summary": "linked",
            }]}
            write_json_atomic(planning / "navigation.json", {"pages": {"001.md": state}})
            write_json_atomic(peer_planning / "navigation.json", {"pages": {"001.md": state}})

            class Model:
                def __init__(self):
                    self.page_plan_calls = 0

                async def structured(self, schema, _messages, **_kwargs):
                    if schema is EdgeSuggestions:
                        return EdgeSuggestions(edges=[{
                            "target_node_id": peer.chunk_id,
                            "label": "uses", "summary": "linked",
                        }])
                    if schema is PageReferencePlan:
                        self.page_plan_calls += 1
                        raise AssertionError("unchanged candidate pool should reuse the current page plan")
                    raise AssertionError(schema)

            model = Model()
            settings = SimpleNamespace(
                wiki_linker_mode="legacy", wiki_output_language="日本語",
                wiki_linker_concurrency=1, concurrency=1,
            )
            result = asyncio.run(link_document(
                project, rel, model=model, embedder=None, settings=settings,
                changed_pages={"doc/001.md"}, render=True,
            ))

            self.assertIn("peer/001.md", result.affected_pages or [])
            self.assertEqual(model.page_plan_calls, 0)
            self.assertIn("New", (project.wiki / "peer" / "001.md").read_text(encoding="utf-8"))

    def test_candidate_copies_the_live_linker_database(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp) / "project"
            mount = Path(tmp) / "mount"
            mount.mkdir()
            project = Project(root, mount).ensure()
            ensure_repository(project)
            connection = sqlite3.connect(project.linker_database)
            try:
                connection.execute("CREATE TABLE proof(value TEXT)")
                connection.execute("INSERT INTO proof VALUES('kept')")
                connection.commit()
            finally:
                connection.close()

            with candidate(project, "copy-linker") as staged:
                connection = sqlite3.connect(staged.linker_database)
                try:
                    value = connection.execute("SELECT value FROM proof").fetchone()[0]
                finally:
                    connection.close()

            self.assertEqual(value, "kept")

    def test_history_rejects_generated_data_without_a_ledger(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            mount = root / "mount"
            mount.mkdir()
            project = Project(root, mount).ensure()
            (project.raw / "orphan.md").write_text("not published", encoding="utf-8")

            with self.assertRaisesRegex(RuntimeError, "without pipeline.json"):
                ensure_repository(project)

    def test_history_rejects_orphan_data_with_an_empty_ledger(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            mount = root / "mount"
            mount.mkdir()
            project = Project(root, mount).ensure()
            save_ledger(project.metadata / "pipeline.json", Ledger({}, {}))
            (project.raw / "orphan.md").write_text("not published", encoding="utf-8")

            with self.assertRaisesRegex(RuntimeError, "not represented by pipeline.json"):
                ensure_repository(project)

    def test_queued_blob_survives_git_garbage_collection(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            mount = root / "mount"
            mount.mkdir()
            source = mount / "a.docx"
            source.write_bytes(b"queued source")
            project = Project(root / "data", mount).ensure()

            blob = stage_blob(project, source)
            subprocess.run(
                ["git", "-C", str(project.root), "gc", "--prune=now"],
                check=True,
                capture_output=True,
            )

            self.assertEqual(read_blob(project, blob.oid), b"queued source")

    def test_page_ownership_marker_survives_a_local_move(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            mount = root / "mount"
            mount.mkdir()
            project = Project(root, mount).ensure()
            old_rel = "team/original_docx.md"
            new_rel = "archive/renamed_docx.md"
            old_folder = project.wiki_dir(old_rel)
            (old_folder / "_planning").mkdir(parents=True)
            (old_folder / "_planning" / "source.json").write_text(
                json.dumps({"raw": old_rel, "id_seed": "team/original.docx"}),
                encoding="utf-8",
            )
            (old_folder / "001.md").write_text("# Page\n\nBody\n", encoding="utf-8")
            publisher = GrowiPublisher(
                object(),
                SimpleNamespace(write_path="/diff", root_path="/diff", mode="attach"),
            )
            old_body = publisher._document_pages(project, old_rel)[0]["body"]
            new_folder = project.wiki_dir(new_rel)
            new_folder.parent.mkdir(parents=True)
            shutil.move(old_folder, new_folder)
            marker_path = new_folder / "_planning" / "source.json"
            marker = json.loads(marker_path.read_text(encoding="utf-8"))
            marker["raw"] = new_rel
            marker_path.write_text(json.dumps(marker), encoding="utf-8")

            new_body = publisher._document_pages(project, new_rel)[0]["body"]
            pattern = re.compile(r"<!-- chunk: ([^ ]+)")
            self.assertEqual(pattern.search(old_body).group(1), pattern.search(new_body).group(1))

    def test_candidate_worker_respects_the_live_pipeline_lock(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            mount = root / "mount"
            mount.mkdir()
            (mount / "a.md").write_text("A", encoding="utf-8")
            settings = SimpleNamespace(data_root=str(root / "data"), target_name="test", mount_path=str(mount))
            project = open_project(settings)
            queue.scan(settings, settle_seconds=0)

            with pipeline._lock(project):
                with self.assertRaisesRegex(RuntimeError, "publisher already running"):
                    queue.work_once(settings)

    def test_publish_rollback_checkpoints_the_restored_revision(self) -> None:
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp)
            mount = base / "mount"
            mount.mkdir()
            source = mount / "test.docx"
            source.write_bytes(b"source")
            project = Project(base / "live", mount).ensure()
            raw_rel = "test_docx.md"
            raw = project.raw_file(raw_rel)
            raw.write_text("# Raw\n", encoding="utf-8")
            wiki = project.wiki_dir(raw_rel)
            (wiki / "_planning").mkdir(parents=True)
            (wiki / "001.md").write_text("# Page\n", encoding="utf-8")
            (wiki / "_planning" / "source.json").write_text(
                json.dumps({
                    "raw": raw_rel,
                    "sha256": hashlib.sha256(raw.read_bytes()).hexdigest(),
                    "id_seed": "test.docx",
                }),
                encoding="utf-8",
            )
            document = wiki.relative_to(project.wiki).as_posix()
            ledger = Ledger(
                {"test.docx": {
                    "source_id": "source-1",
                    "id_seed": "test.docx",
                    "raw_rel": raw_rel,
                    "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
                }},
                {document: {
                    "content_sha256": pipeline._content_hash(wiki),
                    "growi_path": f"/target/{document}",
                    "raw_rel": raw_rel,
                }},
                {f"{document}/001.md": {
                    "growi_path": f"/target/{document}/001",
                    "page_id": "page-1",
                    "revision_id": "base-revision",
                }},
            )
            save_ledger(project.metadata / "pipeline.json", ledger)
            before = ensure_repository(project)

            candidate = Project(base / "candidate", base / "candidate-mount").ensure()
            shutil.copytree(project.raw, candidate.raw, dirs_exist_ok=True)
            shutil.copytree(project.wiki, candidate.wiki, dirs_exist_ok=True)
            candidate_ledger = load_ledger(project.metadata / "pipeline.json")
            candidate_ledger.published_documents[document]["content_sha256"] = "candidate-content"
            candidate_ledger.published_pages[f"{document}/001.md"]["revision_id"] = "candidate-revision"
            save_ledger(candidate.metadata / "pipeline.json", candidate_ledger)

            class Client:
                async def get_page(self, *, page_id=None, path=None):
                    return GrowiPage(
                        page_id="page-1",
                        revision_id="candidate-revision",
                        path=f"/target/{document}/001",
                        body="",
                    )

                async def list_all_pages(self, _path):
                    return []

            class Publisher:
                client = Client()

                def doc_path(self, _project, _rel):
                    return f"/target/{document}"

                def page_marker_id(self, _project, _local_path):
                    return "stable-marker"

                def publish_documents(self, _project, _rels, known_pages=None):
                    return {f"{document}/001.md": GrowiPage(
                        page_id="page-1",
                        revision_id="restored-revision",
                        path=f"/target/{document}/001",
                        body="",
                    )}

                def delete_document(self, _project, _rel):
                    return 0

            settings = SimpleNamespace(
                data_root=str(base), target_name="live", mount_path=str(mount)
            )
            with patch.object(pipeline, "_publisher", return_value=Publisher()):
                restored = pipeline.restore_publication(
                    settings,
                    candidate,
                    [],
                    known_revisions={"page-1": {"candidate-revision"}},
                )

            self.assertNotEqual(before, restored)
            self.assertEqual(restored, last_good(project))
            page = load_ledger(project.metadata / "pipeline.json").published_pages[f"{document}/001.md"]
            self.assertEqual(page["revision_id"], "restored-revision")


@unittest.skipUnless(RUN_LIVE, "set RUN_DIFF_INTEGRATION=1 to run live DOCX diff tests")
class MountDiffPipelineAcceptanceTest(unittest.TestCase):
    """Black-box contract for the implementation described in plan_diff.md."""

    maxDiff = None

    def setUp(self) -> None:
        if shutil.which("git") is None:
            self.skipTest("git is required")
        config.PROJECT_ROOT = ROOT
        base = Settings.from_env(str(CONFIG))
        if not base.growi_url or not base.growi_token:
            self.skipTest("GROWI_URL and GROWI_TOKEN are required")
        if not base.parser_base_url:
            self.skipTest("the DOCX parser URL is required")

        self._tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self._tmp.cleanup)
        sandbox = Path(self._tmp.name)
        mount = sandbox / "mount"
        mount.mkdir()
        self.source = mount / "test.docx"
        shutil.copy2(FIXTURE, self.source)
        suffix = uuid.uuid4().hex[:10]
        target = f"diff-test-{suffix}"
        self.settings = base.model_copy(
            update={
                "data_root": str(sandbox / "data"),
                "mount_path": str(mount),
                "target_name": target,
                "concurrency": 1,
                "wiki_rewrite_concurrency": 1,
                "wiki_linker_concurrency": 1,
                "ingest_concurrency": 1,
                "service_max_agents": 1,
            }
        )
        self.project = open_project(self.settings)
        self.raw_rel = raw_name_for(self.source.name)
        self.events: list[dict] = []
        self.remote_may_exist = False
        self.addCleanup(self._cleanup_remote)

    def _cleanup_remote(self) -> None:
        if not self.remote_may_exist:
            return
        try:
            pipeline.reset_growi(self.settings)
        except Exception as exc:  # cleanup must not hide the test's real failure
            print(f"warning: failed to clean {self.settings.target_name}: {exc}")

    def _scan(self, **kwargs):
        result = queue.scan(self.settings, settle_seconds=0, **kwargs)
        return result

    def _work(self, callback=None):
        events = self.events.append if callback is None else callback
        result = queue.work_once(self.settings, on_event=events)
        if result and not result.get("failures") and any(
            row.get("status") in {"added", "changed", "deleted", "moved"}
            for row in result.get("done", [])
        ):
            self.remote_may_exist = True
        return result

    def _drain(self, limit: int = 8) -> list[dict]:
        results: list[dict] = []
        for _ in range(limit):
            if not queue.status(self.project):
                return results
            result = self._work()
            self.assertIsNotNone(result, f"queue stopped with rows: {queue.status(self.project)}")
            self.assertFalse(result.get("failures"), result)
            results.append(result)
        self.fail(f"queue did not drain after {limit} jobs: {queue.status(self.project)}")

    def _build_initial(self):
        scan = self._scan()
        self.assertEqual(scan["added"], ["test.docx"])
        result = self._work()
        self.assertIsNotNone(result)
        self.assertFalse(result["cancelled"], result)
        self.assertFalse(result["failures"], result)
        self.remote_may_exist = True
        self._assert_built("test.docx")
        return result

    def _git(self, *args: str) -> str:
        result = subprocess.run(
            ["git", "-C", str(self.project.root), *args],
            check=True,
            capture_output=True,
            text=True,
        )
        return result.stdout.strip()

    def _last_good(self) -> str:
        return self._git("rev-parse", LAST_GOOD_REF)

    def _commit_count(self) -> int:
        return int(self._git("rev-list", "--count", LAST_GOOD_REF))

    def _ledger(self):
        return load_ledger(self.project.metadata / "pipeline.json")

    def _source_row(self, rel: str = "test.docx") -> dict:
        row = self._ledger().sources.get(rel)
        self.assertIsNotNone(row, f"missing ledger source {rel}")
        return row

    def _assert_built(self, rel: str) -> None:
        row = self._source_row(rel)
        raw_rel = str(row["raw_rel"])
        project_path = self.project.raw_file(raw_rel)
        self.assertTrue(project_path.is_file(), project_path)
        self.assertTrue(self.project.wiki_dir(raw_rel).is_dir())
        document = self.project.wiki_dir(raw_rel).relative_to(self.project.wiki).as_posix()
        self.assertIn(document, self._ledger().published_documents)
        self.assertTrue(row.get("source_id"), row)
        self.assertTrue(row.get("source_blob_oid"), row)

    def _done(self, result: dict, rel: str) -> dict:
        return next(row for row in result["done"] if row.get("path") == rel)

    def _inject_once(self, stage: str, action, *, step: str = "start"):
        fired = False

        def callback(event: dict) -> None:
            nonlocal fired
            self.events.append(dict(event))
            if not fired and event.get("stage") == stage and event.get("step") == step:
                fired = True
                action()

        return callback

    def _published_page_ids(self) -> set[str]:
        return {
            str(row["page_id"])
            for row in self._ledger().published_pages.values()
            if row.get("page_id")
        }

    def _chunk_ids(self, rel: str) -> set[str]:
        raw_rel = str(self._source_row(rel)["raw_rel"])
        data = json.loads((self.project.wiki_dir(raw_rel) / "_planning" / "chunks.json").read_text(encoding="utf-8"))
        return {
            str(chunk["chunk_id"])
            for page in data.get("pages", [])
            for chunk in page.get("chunks", [])
        }

    def _remote_pages(self):
        publisher = pipeline._publisher(self.settings)
        self.assertIsNotNone(publisher)
        return asyncio.run(publisher.client.list_all_pages(publisher.connection.root_path))

    def test_00_history_repository_excludes_runtime_state(self) -> None:
        self._scan()
        self.assertTrue((self.project.root / ".git").is_dir())
        self.assertTrue(self._last_good())
        for relative in (
            "mount/test.docx",
            "metadata/watch-queue.sqlite",
            "metadata/watch-queue.sqlite-wal",
            "metadata/wiki-linker.sqlite",
            "metadata/pipeline.lock",
            "metadata/work/example",
            "metadata/candidates/example",
        ):
            ignored = subprocess.run(
                ["git", "-C", str(self.project.root), "check-ignore", "-q", relative]
            )
            self.assertEqual(ignored.returncode, 0, f"runtime path is not ignored: {relative}")

    def test_01_new_file_added_is_published_and_committed(self) -> None:
        self._scan()
        before = self._last_good()
        result = self._work()
        self.assertFalse(result["failures"], result)
        self.assertNotEqual(before, self._last_good())
        self.assertEqual(self._done(result, "test.docx")["rebuild"], "full")
        self._assert_built("test.docx")

    def test_01b_committed_file_deleted_is_removed_and_committed(self) -> None:
        self._build_initial()
        before = self._last_good()
        old_ids = self._published_page_ids()
        self.source.unlink()
        self.assertEqual(self._scan()["deleted"], ["test.docx"])
        result = self._work()
        self.assertFalse(result["failures"], result)
        self.assertNotEqual(before, self._last_good())
        self.assertNotIn("test.docx", self._ledger().sources)
        self.assertFalse(self.project.raw_file(self.raw_rel).exists())
        self.assertFalse(self.project.wiki_dir(self.raw_rel).exists())
        remote_ids = {page.page_id for page in self._remote_pages()}
        self.assertTrue(old_ids.isdisjoint(remote_ids))

    def test_02_add_then_delete_before_first_scan_is_a_noop(self) -> None:
        self.source.unlink()
        result = self._scan()
        self.assertEqual(result["added"], [])
        self.assertEqual(result["deleted"], [])
        self.assertEqual(queue.status(self.project), [])
        self.assertIsNone(self._work())
        self.assertFalse(self.project.raw_file(self.raw_rel).exists())

    def test_03_add_then_delete_while_queued_is_cancelled(self) -> None:
        self._scan()
        before = self._last_good()
        self.source.unlink()
        result = self._scan()
        self.assertEqual(result["cancelled"], ["test.docx"])
        self.assertEqual(queue.status(self.project), [])
        self.assertEqual(before, self._last_good())
        self.assertIsNone(self._work())

    def test_04a_new_file_deleted_during_generation_discards_candidate(self) -> None:
        self._scan()
        before = self._last_good()

        def remove_source() -> None:
            self.source.unlink()
            self.assertEqual(self._scan()["cancelled"], ["test.docx"])

        result = self._work(self._inject_once("wiki", remove_source))
        self.assertTrue(result["cancelled"], result)
        self.assertEqual(before, self._last_good())
        self.assertEqual(queue.status(self.project), [])
        self.assertFalse(self.project.raw_file(self.raw_rel).exists())
        self.assertFalse(self.project.wiki_dir(self.raw_rel).exists())
        self.assertNotIn("test.docx", self._ledger().sources)

    def test_04b_committed_file_deleted_during_generation_commits_delete(self) -> None:
        self._build_initial()
        make_major_docx_change(self.source, "delete-during-build")
        self._scan()
        before = self._last_good()

        def remove_source() -> None:
            self.source.unlink()
            scan = self._scan()
            self.assertEqual(scan["deleted"], ["test.docx"])

        first = self._work(self._inject_once("wiki", remove_source))
        self.assertTrue(first["cancelled"], first)
        self.assertEqual(before, self._last_good())
        self.assertEqual(queue.status(self.project)[0]["operation"], "delete")
        self._drain()
        self.assertNotEqual(before, self._last_good())
        self.assertNotIn("test.docx", self._ledger().sources)

    def test_05_add_then_update_before_first_scan_builds_latest_only(self) -> None:
        marker = "case-five-latest"
        make_minor_docx_change(self.source, marker)
        self.assertEqual(self._scan()["added"], ["test.docx"])
        result = self._work()
        self.assertFalse(result["failures"], result)
        raw = self.project.raw_file(self.raw_rel).read_text(encoding="utf-8")
        self.assertIn(marker, raw)
        self.assertEqual(self._commit_count(), 2)  # initial baseline plus one add

    def test_06_add_then_update_while_queued_replaces_target(self) -> None:
        first = self._scan()
        first_sha = queue.status(self.project)[0]["target_sha256"]
        marker = "case-six-newest"
        make_minor_docx_change(self.source, marker)
        second = self._scan()
        rows = queue.status(self.project)
        self.assertEqual(len(rows), 1)
        self.assertNotEqual(first_sha, rows[0]["target_sha256"])
        self.assertEqual(first["added"], ["test.docx"])
        self.assertEqual(second["updated"], ["test.docx"])
        result = self._work()
        self.assertFalse(result["failures"], result)
        self.assertIn(marker, self.project.raw_file(self.raw_rel).read_text(encoding="utf-8"))
        self.assertEqual(self._commit_count(), 2)

    def test_07a_small_update_during_generation_finishes_then_runs_incrementally(self) -> None:
        self._scan()
        initial = self._last_good()
        scans: list[dict] = []

        def update_source() -> None:
            make_minor_docx_change(self.source, "case-seven-small")
            scans.append(self._scan())

        first = self._work(self._inject_once("wiki", update_source))
        self.assertFalse(first["cancelled"], first)
        after_first = self._last_good()
        self.assertNotEqual(initial, after_first)
        self.assertEqual(scans[0]["classification"]["test.docx"]["kind"], "small")
        self.assertEqual(len(queue.status(self.project)), 1)
        second = self._work()
        self.assertFalse(second["cancelled"], second)
        self.assertEqual(self._done(second, "test.docx")["rebuild"], "incremental")
        self.assertNotEqual(after_first, self._last_good())
        self.assertIn("case-seven-small", self.project.raw_file(self.raw_rel).read_text(encoding="utf-8"))

    def test_07b_large_update_during_generation_cancels_and_restarts_from_last_good(self) -> None:
        self._scan()
        initial = self._last_good()
        scans: list[dict] = []

        def update_source() -> None:
            make_major_docx_change(self.source, "case-seven-large")
            scans.append(self._scan())

        first = self._work(self._inject_once("wiki", update_source))
        self.assertTrue(first["cancelled"], first)
        self.assertEqual(initial, self._last_good())
        self.assertEqual(scans[0]["classification"]["test.docx"]["kind"], "large")
        second = self._work()
        self.assertFalse(second["cancelled"], second)
        self.assertEqual(self._done(second, "test.docx")["rebuild"], "full")
        self.assertEqual(self._commit_count(), 2)  # cancelled candidate was never promoted

    def test_08a_small_update_after_completion_is_incremental(self) -> None:
        self._build_initial()
        before = self._last_good()
        make_minor_docx_change(self.source, "case-eight-small")
        scan = self._scan()
        self.assertEqual(scan["classification"]["test.docx"]["kind"], "small")
        result = self._work()
        self.assertEqual(self._done(result, "test.docx")["rebuild"], "incremental")
        self.assertNotEqual(before, self._last_good())

    def test_08b_large_update_after_completion_is_full(self) -> None:
        self._build_initial()
        make_major_docx_change(self.source, "case-eight-large")
        scan = self._scan()
        self.assertEqual(scan["classification"]["test.docx"]["kind"], "large")
        result = self._work()
        self.assertEqual(self._done(result, "test.docx")["rebuild"], "full")

    def test_09_rename_preserves_source_chunk_and_growi_page_ids(self) -> None:
        self._build_initial()
        old_source = self._source_row()
        old_source_id = old_source["source_id"]
        old_seed = old_source["id_seed"]
        old_chunks = self._chunk_ids("test.docx")
        old_pages = self._published_page_ids()
        destination = self.source.parent / "renamed" / "test.docx"
        destination.parent.mkdir()
        self.source.rename(destination)
        scan = self._scan()
        self.assertEqual(scan["moved"], [{"from": "test.docx", "to": "renamed/test.docx"}])
        row = queue.status(self.project)[0]
        self.assertEqual(row["operation"], "move")
        self.assertEqual(row["from_rel"], "test.docx")
        result = self._work()
        self.assertEqual(self._done(result, "renamed/test.docx")["rebuild"], "move")
        new_source = self._source_row("renamed/test.docx")
        self.assertEqual(new_source["source_id"], old_source_id)
        self.assertEqual(new_source["id_seed"], old_seed)
        self.assertEqual(self._chunk_ids("renamed/test.docx"), old_chunks)
        self.assertEqual(self._published_page_ids(), old_pages)
        self.assertNotIn("test.docx", self._ledger().sources)

    def test_10_delete_then_readd_before_work_keeps_identity_and_latest_content(self) -> None:
        self._build_initial()
        source_id = self._source_row()["source_id"]
        self.source.unlink()
        self.assertEqual(self._scan()["deleted"], ["test.docx"])
        shutil.copy2(FIXTURE, self.source)
        make_minor_docx_change(self.source, "delete-readd-final")
        self._scan()
        rows = queue.status(self.project)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["operation"], "update")
        result = self._work()
        self.assertFalse(result["failures"], result)
        self.assertEqual(self._source_row()["source_id"], source_id)
        self.assertIn("delete-readd-final", self.project.raw_file(self.raw_rel).read_text(encoding="utf-8"))

    def test_11_repeated_rapid_updates_keep_one_latest_target(self) -> None:
        self._build_initial()
        for number in range(1, 4):
            make_minor_docx_change(self.source, f"rapid-{number}")
            self._scan()
        rows = queue.status(self.project)
        self.assertEqual(len(rows), 1)
        target_sha = rows[0]["target_sha256"]
        result = self._work()
        self.assertFalse(result["failures"], result)
        raw = self.project.raw_file(self.raw_rel).read_text(encoding="utf-8")
        self.assertIn("rapid-3", raw)
        self.assertEqual(self._source_row()["source_sha256"], target_sha)

    def test_12_worker_restart_recovers_claimed_job(self) -> None:
        self._scan()
        claimed = queue.claim(self.project, "slow")
        self.assertEqual(len(claimed), 1)
        self.assertEqual(queue.status(self.project)[0]["status"], "running")
        self.assertEqual(queue.recover(self.project), 1)
        self.assertEqual(queue.status(self.project)[0]["status"], "queued")
        result = self._work()
        self.assertFalse(result["failures"], result)
        self._assert_built("test.docx")

    def test_13_parse_failure_keeps_last_good_unchanged(self) -> None:
        self._scan()
        before = self._last_good()
        with patch.object(pipeline, "_parse", side_effect=RuntimeError("injected parse failure")):
            result = self._work()
        self.assertTrue(result["failures"], result)
        self.assertEqual(before, self._last_good())
        self.assertFalse(self.project.raw_file(self.raw_rel).exists())
        self.assertFalse(self.project.wiki_dir(self.raw_rel).exists())
        self.assertEqual(queue.status(self.project)[0]["status"], "failed")

    def test_14_generation_failure_keeps_last_good_unchanged(self) -> None:
        self._scan()
        before = self._last_good()
        with patch.object(pipeline, "write_wiki_pages", side_effect=RuntimeError("injected generation failure")):
            result = self._work()
        self.assertTrue(result["failures"], result)
        self.assertEqual(before, self._last_good())
        self.assertFalse(self.project.raw_file(self.raw_rel).exists())
        self.assertFalse(self.project.wiki_dir(self.raw_rel).exists())

    def test_15_publish_failure_restores_last_good_locally_and_remotely(self) -> None:
        self._scan()
        before = self._last_good()
        original = pipeline.GrowiPublisher.publish_documents

        def publish_then_fail(publisher, *args, **kwargs):
            original(publisher, *args, **kwargs)
            self.remote_may_exist = True
            raise RuntimeError("injected failure after remote writes")

        with patch.object(pipeline.GrowiPublisher, "publish_documents", new=publish_then_fail):
            result = self._work()
        self.assertTrue(result["failures"], result)
        self.assertEqual(before, self._last_good())
        self.assertFalse(self.project.raw_file(self.raw_rel).exists())
        self.assertFalse(self.project.wiki_dir(self.raw_rel).exists())
        remote = [page.path for page in self._remote_pages() if page.path.rstrip("/") != f"/{self.settings.target_name}"]
        self.assertEqual(remote, [])

    def test_16_content_audit_detects_change_with_unchanged_metadata(self) -> None:
        self._build_initial()
        old_snapshot = queue._snapshot(self.project.mount)
        make_minor_docx_change(self.source, "same-metadata-audit")
        with patch.object(queue, "_snapshot", return_value=old_snapshot):
            result = self._scan(verify_content=True)
        self.assertEqual(result["updated"], ["test.docx"])
        self.assertEqual(result["classification"]["test.docx"]["kind"], "small")

    def test_17_small_update_during_linking_finishes_then_runs(self) -> None:
        self._scan()

        def update_source() -> None:
            make_minor_docx_change(self.source, "update-during-link")
            self._scan()

        first = self._work(self._inject_once("linker", update_source, step="batch_start"))
        self.assertFalse(first["cancelled"], first)
        self.assertEqual(len(queue.status(self.project)), 1)
        second = self._work()
        self.assertFalse(second["cancelled"], second)
        self.assertIn("update-during-link", self.project.raw_file(self.raw_rel).read_text(encoding="utf-8"))

    def test_18_update_during_publish_finishes_current_commit_then_runs_next(self) -> None:
        self._scan()
        scans: list[dict] = []

        def update_source() -> None:
            make_major_docx_change(self.source, "update-during-publish")
            scans.append(self._scan())

        first = self._work(self._inject_once("growi-publish", update_source))
        self.assertFalse(first["cancelled"], first)
        first_commit = self._last_good()
        self.assertEqual(scans[0]["classification"]["test.docx"]["kind"], "large")
        self.assertEqual(len(queue.status(self.project)), 1)
        second = self._work()
        self.assertFalse(second["cancelled"], second)
        self.assertNotEqual(first_commit, self._last_good())
        self.assertEqual(self._done(second, "test.docx")["rebuild"], "full")

    def test_19_forced_sync_rebuilds_unchanged_source(self) -> None:
        self._build_initial()
        scan = self._scan(force=True)
        self.assertEqual(scan["classification"]["test.docx"]["kind"], "forced")
        result = self._work()
        done = self._done(result, "test.docx")
        self.assertEqual((done["tier"], done["reason"]), (3, "forced"))

    def test_20_small_edit_stays_on_the_fast_path(self) -> None:
        self._build_initial()
        self.events.clear()
        make_minor_docx_change(self.source, "case-twenty-small")
        self._scan()

        def record(event: dict) -> None:
            self.events.append(dict(event))

        result = self._work(record)
        done = self._done(result, "test.docx")
        self.assertEqual(done["tier"], 1)
        self.assertFalse(any(event.get("stage") == "seed" and event.get("step") == "start" for event in self.events))
        self.assertFalse(any(event.get("stage") == "rewrite" and event.get("step") == "page_done" for event in self.events))
        decision = next(
            event for event in reversed(self.events)
            if event.get("stage") == "wiki" and event.get("step") == "update_decision" and event.get("file") == self.raw_rel
        )
        changed_pages = set(decision.get("changed_pages", []))
        curated = sum(event.get("stage") == "linker" and event.get("step") == "page_curated" for event in self.events)
        self.assertLessEqual(
            curated, len(changed_pages),
            f"curated {curated} pages for {len(changed_pages)} changed pages: {sorted(changed_pages)}",
        )

    def test_21_section_rewrite_regenerates_only_that_section(self) -> None:
        self._build_initial()
        self.events.clear()
        document = Document(self.source)
        paragraphs = document.paragraphs
        heading_pattern = re.compile(r"^[0-9０-９]+[．.、)）]")
        headings = [index for index, paragraph in enumerate(paragraphs) if heading_pattern.match(paragraph.text.strip())]
        self.assertGreaterEqual(len(headings), 4, "Japanese fixture needs at least four numbered sections")
        start = headings[2] + 1
        end = headings[3]
        changed = [paragraph for paragraph in paragraphs[start:end] if paragraph.text.strip()]
        self.assertTrue(changed, "third numbered section should contain paragraphs")
        for index, paragraph in enumerate(changed, 1):
            paragraph.text = f"改訂確認 第三節 {index} LLMWIKI-SECTION-UPDATE"
        document.save(self.source)
        self._scan()
        result = self._work()
        done = self._done(result, "test.docx")
        self.assertEqual(done["tier"], 2)
        decision = next(
            event for event in reversed(self.events)
            if event.get("stage") == "wiki" and event.get("step") == "update_decision" and event.get("file") == self.raw_rel
        )
        regenerated = int(decision["regenerate_pages"])
        page_done = sum(event.get("stage") == "rewrite" and event.get("step") == "page_done" for event in self.events)
        page_count = len(list((self.project.state_dir(self.raw_rel) / "state" / "pages").glob("*.json")))
        self.assertEqual(page_done, regenerated)
        self.assertLess(regenerated, page_count / 2, f"regenerated {regenerated} of {page_count} pages")

    @unittest.skipUnless(shutil.which("soffice") or shutil.which("libreoffice"), "LibreOffice is required")
    def test_22_legacy_doc_is_converted_and_published(self) -> None:
        executable = shutil.which("soffice") or shutil.which("libreoffice")
        with tempfile.TemporaryDirectory(prefix="legacy-doc-profile-") as profile:
            subprocess.run([
                executable, "--headless", "--nologo", "--nodefault", "--nolockcheck", "--nofirststartwizard",
                f"-env:UserInstallation={Path(profile).as_uri()}",
                "--convert-to", "doc", "--outdir", str(self.source.parent), str(self.source),
            ], check=True, capture_output=True, timeout=120)
        self.source.unlink()
        self._scan()
        self._drain()
        row = self._source_row("test.doc")
        self.assertEqual(row["parser"], "doc")
        self.assertTrue(self.project.wiki_dir(raw_name_for("test.doc")).is_dir())


if __name__ == "__main__":
    unittest.main()
