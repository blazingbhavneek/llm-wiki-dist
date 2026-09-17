"""Cross-deployment equivalence for the XLSM lineage builder.

The graph and parser deployments ship independent copies of the static
lineage analysis (see plan §25); this test locks their public schemas so the
generic final-output pages and the llm-wiki manifest stay aligned even as the
two services evolve independently.
"""

from __future__ import annotations

import importlib.util
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from graph.workspace import xlsm as graph_lineage

_PARSER_MODULE_PATH = (
    Path(__file__).resolve().parents[1]
    .parent
    / "parser"
    / "formats"
    / "xlsm_lineage.py"
)


def _load_parser_lineage():
    spec = importlib.util.spec_from_file_location(
        "parser_xlsm_lineage", _PARSER_MODULE_PATH
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _fixture_workbook(path: Path) -> None:
    workbook = Workbook()
    inputs = workbook.active
    inputs.title = "Inputs"
    inputs["A1"] = 1
    inputs["A2"] = 2
    outputs = workbook.create_sheet("Outputs")
    outputs["B1"] = "=SUM(Inputs!A1:A2)"
    hidden = workbook.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["B1"] = "=Outputs!B1"
    workbook.save(path)
    workbook.close()


_MODULES = [
    (
        "Module1.bas",
        "Public Const TARGET As String = \"Outputs\"\n"
        "Sub Update()\n"
        "Dim ws As Worksheet\n"
        "Set ws = Worksheets(TARGET)\n"
        "ws.Range(\"B2\").ClearContents\n"
        "ws.Range(\"B1\").Value = 99\n"
        "End Sub\n",
    )
]


class LineageEquivalenceTests(unittest.TestCase):
    def test_parser_lineage_module_exists(self) -> None:
        self.assertTrue(_PARSER_MODULE_PATH.is_file(), _PARSER_MODULE_PATH)

    def test_build_manifests_agree(self) -> None:
        parser_lineage = _load_parser_lineage()

        # Force has_vba() to accept the fixture workbook, which openpyxl cannot
        # create with a real VBA project in this offline test.
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "sample.xlsm"
            _fixture_workbook(path)
            with (
                patch.object(graph_lineage, "has_vba", lambda _p: True),
                patch.object(parser_lineage, "has_vba", lambda _p: True),
                patch.object(graph_lineage, "_modules", lambda _p: list(_MODULES)),
                patch.object(parser_lineage, "_modules", lambda _p: list(_MODULES)),
            ):
                graph_manifest = graph_lineage.build_manifest(path)
                parser_manifest = parser_lineage.build_manifest(path)

        self.assertIsNotNone(graph_manifest)
        self.assertEqual(graph_manifest, parser_manifest)

        # Both sides must expose the exact schema fields plan §25 requires.
        for key in (
            "schema_version",
            "mode",
            "sheets",
            "procedures",
        ):
            self.assertIn(key, graph_manifest)
        for procedure in graph_manifest["procedures"]:
            for key in (
                "final_outputs_affected",
                "buttons",
                "formula_cells",
                "sheets_read",
                "sheets_written",
                "unresolved_dynamic_references",
            ):
                self.assertIn(key, procedure)

    def test_generic_render_helpers_present(self) -> None:
        parser_lineage = _load_parser_lineage()
        self.assertTrue(hasattr(parser_lineage, "render_consolidated_vba_code"))
        self.assertTrue(
            hasattr(parser_lineage, "render_consolidated_vba_final_outputs")
        )


if __name__ == "__main__":
    unittest.main()
