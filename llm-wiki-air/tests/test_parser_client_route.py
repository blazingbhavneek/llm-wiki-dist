from __future__ import annotations

import io
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

import graph.workspace.parser_client as parser_client
from graph.wiki.images import neutralize_image_descriptions, reuse_image_descriptions
from graph.workspace.parser_client import LLM_WIKI_PARSE_PATH, parse_document


class _Resp:
    def __init__(self, payload: dict) -> None:
        self._payload = payload
        self.status_code = 200
        self.text = ""

    def json(self) -> dict:
        return self._payload

    def raise_for_status(self) -> None:
        return None


class _Settings:
    chat_base_url = "http://llm/v1"
    chat_api_key = "k"
    chat_model = "m"


def _image(payload: str, description: str = "") -> str:
    return (
        '<image-unit>\n  <image-media><img src="data:image/png;base64,'
        f'{payload}" alt="diagram"></image-media>\n'
        f"  <image-description>{description}</image-description>\n"
        "</image-unit>"
    )


class ParserClientTests(unittest.TestCase):
    def test_uses_the_llm_wiki_route(self) -> None:
        self.assertEqual(LLM_WIKI_PARSE_PATH, "/parse/llm-wiki")

        stream = io.BytesIO()
        workbook = Workbook()
        workbook.active["A1"] = 1
        workbook.save(stream)
        workbook.close()

        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            path.write_bytes(stream.getvalue())
            with patch.object(
                parser_client.requests,
                "post",
                return_value=_Resp({"markdown": "# x", "pages": ["## foo"]}),
            ) as post:
                result = parse_document(
                    path,
                    base_url="http://parser/agent/doc-parser",
                    settings=_Settings(),
                )

            self.assertEqual(result, "# x")
        self.assertEqual(
            post.call_args.args[0],
            "http://parser/agent/doc-parser/parse/llm-wiki",
        )

    def test_rejects_malformed_pages(self) -> None:
        stream = io.BytesIO()
        Workbook().save(stream)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            path.write_bytes(stream.getvalue())
            with patch.object(
                parser_client.requests,
                "post",
                return_value=_Resp({"markdown": "# x", "pages": "not-a-list"}),
            ):
                with self.assertRaisesRegex(RuntimeError, "pages must be a list"):
                    parse_document(
                        path,
                        base_url="http://parser",
                        settings=_Settings(),
                    )

    def test_update_reuses_unchanged_image_description_without_vision_call(self) -> None:
        previous = _image("YWJj", "old description")
        current = _image("YWJj")
        stream = io.BytesIO()
        Workbook().save(stream)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            path.write_bytes(stream.getvalue())
            with patch.object(
                parser_client.requests,
                "post",
                return_value=_Resp({"markdown": current}),
            ) as post:
                result = parse_document(
                    path,
                    base_url="http://parser",
                    settings=_Settings(),
                    previous_markdown=previous,
                )

        self.assertIn("<image-description>old description</image-description>", result)
        self.assertEqual(post.call_count, 1)
        self.assertEqual(post.call_args.kwargs["params"]["describe_images"], "false")

    def test_update_describes_only_new_image_bytes(self) -> None:
        previous = _image("YWJj", "keep me")
        current = previous.replace("keep me", "") + "\n" + _image("ZGVm")
        stream = io.BytesIO()
        Workbook().save(stream)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "input.xlsx"
            path.write_bytes(stream.getvalue())
            with patch.object(
                parser_client.requests,
                "post",
                side_effect=[
                    _Resp({"markdown": current}),
                    _Resp({"choices": [{"message": {"content": "new description"}}]}),
                ],
            ) as post:
                result = parse_document(
                    path,
                    base_url="http://parser",
                    settings=_Settings(),
                    previous_markdown=previous,
                )

        self.assertEqual(post.call_count, 2)
        self.assertIn("<image-description>keep me</image-description>", result)
        self.assertIn("<image-description>new description</image-description>", result)
        self.assertEqual(post.call_args_list[1].args[0], "http://llm/v1/chat/completions")

    def test_image_delete_needs_no_description_call(self) -> None:
        calls: list[str] = []
        result = reuse_image_descriptions(
            "before\n" + _image("YWJj", "old") + "\nafter",
            "before\nafter",
            lambda data_url, alt: calls.append(data_url) or "unexpected",
        )
        self.assertEqual(result, "before\nafter")
        self.assertEqual(calls, [])

    def test_description_wording_is_not_part_of_source_diff(self) -> None:
        old = _image("YWJj", "first wording")
        new = _image("YWJj", "completely different wording")
        self.assertEqual(
            neutralize_image_descriptions(old),
            neutralize_image_descriptions(new),
        )


if __name__ == "__main__":
    unittest.main()
