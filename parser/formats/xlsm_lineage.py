"""Static XLSM lineage inspection. Workbook macros are read, never executed.

This module lives in the parser deployment so generic uploads can produce
lineage data without any downstream dependency, and the separate graph
pipeline can keep its own copy (see schema-equivalence test). Both copies
implement the same public ``build_manifest`` schema so parser-side generic
VBA final-output pages and downstream manifest-driven pages agree exactly.
"""

from __future__ import annotations

import posixpath
import re
import zipfile
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

from openpyxl import load_workbook
from openpyxl.formula import Tokenizer
from openpyxl.utils import get_column_letter

_RID = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
_PROC_RE = re.compile(
    r"^\s*(?:(?:Public|Private|Friend|Static)\s+)?"
    r"(Sub|Function|Property\s+(?:Get|Let|Set))\s+([^\s(]+)",
    re.IGNORECASE | re.MULTILINE,
)
_CONST_RE = re.compile(r'^\s*Const\s+(\w+)(?:\s+As\s+\w+)?\s*=\s*"([^"]+)"', re.IGNORECASE | re.MULTILINE)
_SHEET_CALL_RE = re.compile(r"(?:Worksheets|Sheets)\s*\(\s*([^\)]+)\s*\)", re.IGNORECASE)
_ALIAS_RE = re.compile(r"^\s*Set\s+(\w+)\s*=.*?(?:Worksheets|Sheets)\s*\(\s*([^\)]+)\s*\)", re.IGNORECASE)
_MUTATE_RE = re.compile(r"(?:\.Clear\w*\b|\.Delete\b|\.PasteSpecial\b|Destination\s*:=)", re.IGNORECASE)


def has_vba(path: Path) -> bool:
    if Path(path).suffix.lower() != ".xlsm":
        return False
    try:
        with zipfile.ZipFile(path) as archive:
            return "xl/vbaProject.bin" in archive.namelist()
    except (OSError, zipfile.BadZipFile):
        return False


def _sheet_refs(text: str, names: set[str]) -> set[str]:
    refs: set[str] = set()
    try:
        tokens = Tokenizer(text).items if text.startswith("=") else ()
    except Exception:
        tokens = ()
    for token in tokens:
        if token.type != "OPERAND" or token.subtype != "RANGE" or "!" not in token.value:
            continue
        name = token.value.rsplit("!", 1)[0].strip("'").replace("''", "'")
        if name in names:
            refs.add(name)
    for name in names:
        if f"'{name.replace("'", "''")}'!" in text or f"{name}!" in text:
            refs.add(name)
    return refs


def _relations(archive: zipfile.ZipFile, part: str) -> dict[str, str]:
    folder, filename = posixpath.split(part)
    rels_name = posixpath.join(folder, "_rels", f"{filename}.rels")
    if rels_name not in archive.namelist():
        return {}
    root = ElementTree.fromstring(archive.read(rels_name))
    output: dict[str, str] = {}
    for rel in root:
        if rel.attrib.get("TargetMode") == "External" or not rel.attrib.get("Id"):
            continue
        target = rel.attrib.get("Target", "")
        output[rel.attrib["Id"]] = target.lstrip("/") if target.startswith("/") else posixpath.normpath(posixpath.join(folder, target))
    return output


def _local(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _anchor_cells(anchor: Any) -> str:
    points: list[str] = []
    for marker in (node for node in anchor if _local(node.tag) in {"from", "to"}):
        values = {_local(child.tag): child.text for child in marker}
        try:
            points.append(f"{get_column_letter(int(values['col']) + 1)}{int(values['row']) + 1}")
        except (KeyError, TypeError, ValueError):
            continue
    return points[0] if len(points) == 1 else ":".join(points[:2]) if points else "unknown"


def _buttons(path: Path, workbook: Any) -> list[dict[str, str]]:
    output: list[dict[str, str]] = []
    with zipfile.ZipFile(path) as archive:
        workbook_part = "xl/workbook.xml"
        workbook_rels = _relations(archive, workbook_part)
        sheets = ElementTree.fromstring(archive.read(workbook_part))
        for sheet in (node for node in sheets.iter() if _local(node.tag) == "sheet"):
            name = sheet.attrib.get("name", "")
            part = workbook_rels.get(sheet.attrib.get(_RID, ""), "")
            if not part or part not in archive.namelist():
                continue
            sheet_rels = _relations(archive, part)
            sheet_xml = ElementTree.fromstring(archive.read(part))
            for drawing in (node for node in sheet_xml.iter() if _local(node.tag) in {"drawing", "legacyDrawing"}):
                drawing_part = sheet_rels.get(drawing.attrib.get(_RID, ""), "")
                if not drawing_part or drawing_part not in archive.namelist():
                    continue
                root = ElementTree.fromstring(archive.read(drawing_part))
                if _local(drawing.tag) == "drawing":
                    for anchor in (node for node in root.iter() if _local(node.tag) in {"oneCellAnchor", "twoCellAnchor", "absoluteAnchor"}):
                        macro = next((value for node in anchor.iter() for key, value in node.attrib.items() if _local(key) == "macro" and value), "")
                        if macro:
                            output.append({"sheet": name, "cells": _anchor_cells(anchor), "procedure": macro.rsplit("!", 1)[-1]})
                    continue
                for client in (node for node in root.iter() if _local(node.tag) == "ClientData"):
                    values = {_local(child.tag): (child.text or "").strip() for child in client}
                    macro = values.get("FmlaMacro", "").rsplit("!", 1)[-1]
                    try:
                        anchor = [int(value.strip()) for value in values["Anchor"].split(",")]
                        start = f"{get_column_letter(anchor[0] + 1)}{anchor[2] + 1}"
                        end = f"{get_column_letter(anchor[4] + 1)}{anchor[6] + 1}"
                    except (KeyError, ValueError, IndexError):
                        start = end = "unknown"
                    if macro:
                        output.append({"sheet": name, "cells": start if start == end else f"{start}:{end}", "procedure": macro})
    return output


def _modules(path: Path) -> list[tuple[str, str]]:
    try:
        from oletools.olevba import VBA_Parser
    except ImportError as exc:
        raise RuntimeError("oletools is required to inspect macro-enabled workbooks") from exc
    parser = VBA_Parser(str(path))
    try:
        return [(Path(module).name, str(code).replace("\r\n", "\n")) for _, _, module, code in parser.extract_macros()]
    finally:
        parser.close()


def _procedures(modules: list[tuple[str, str]], sheets: set[str], buttons: list[dict[str, str]]) -> list[dict[str, Any]]:
    constants = {name.casefold(): value for _, code in modules for name, value in _CONST_RE.findall(code) if value in sheets}
    procedures: list[dict[str, Any]] = []
    for module_name, code in modules:
        matches = list(_PROC_RE.finditer(code))
        declarations = "\n".join(
            line
            for line in code[: matches[0].start() if matches else len(code)].splitlines()
            if line.strip() and not line.lstrip().startswith("Attribute VB_")
        )
        if declarations:
            procedures.append({
                "id": f"{Path(module_name).stem.casefold()}::declarations",
                "module": module_name,
                "name": "宣言",
                "kind": "declarations",
                "code": declarations,
                "buttons": [],
                "formula_cells": [],
                "sheets_read": [],
                "sheets_written": [],
                "unresolved_dynamic_references": [],
            })
        for index, match in enumerate(matches):
            body = code[match.start(): matches[index + 1].start() if index + 1 < len(matches) else len(code)].rstrip()
            aliases: dict[str, str] = {}
            reads: set[str] = set()
            writes: set[str] = set()
            unresolved: list[str] = []
            for raw_line in body.splitlines():
                line = raw_line.split("'", 1)[0].strip()
                if not line:
                    continue
                if alias := _ALIAS_RE.search(line):
                    key = alias.group(2).strip().strip('"').casefold()
                    if key in constants:
                        aliases[alias.group(1).casefold()] = constants[key]
                    elif key in {name.casefold() for name in sheets}:
                        aliases[alias.group(1).casefold()] = next(name for name in sheets if name.casefold() == key)
                def referenced(fragment: str) -> set[str]:
                    found = _sheet_refs(fragment, sheets)
                    found.update(name for key, name in constants.items() if re.search(rf"\b{re.escape(key)}\b", fragment, re.IGNORECASE))
                    found.update(name for alias_name, name in aliases.items() if re.search(rf"\b{re.escape(alias_name)}\s*\.", fragment, re.IGNORECASE))
                    for value in _SHEET_CALL_RE.findall(fragment):
                        key = value.strip().strip('"').casefold()
                        if key in constants:
                            found.add(constants[key])
                        else:
                            found.update(name for name in sheets if name.casefold() == key)
                    return found

                found = referenced(line)
                dynamic = [value.strip() for value in _SHEET_CALL_RE.findall(line) if value.strip().strip('"').casefold() not in constants and value.strip().strip('"').casefold() not in {name.casefold() for name in sheets}]
                if dynamic:
                    unresolved.append(raw_line.strip())
                if alias:
                    continue
                assignment = re.search(r"(?<![<>=:])=(?!=)", line)
                if assignment and not line.lower().startswith("set "):
                    written = referenced(line[: assignment.start()])
                    writes.update(written)
                    reads.update(found - written)
                elif _MUTATE_RE.search(line):
                    writes.update(found)
                else:
                    reads.update(found)
            name = match.group(2)
            procedure_buttons = [button for button in buttons if button["procedure"].rsplit(".", 1)[-1].casefold() == name.casefold()]
            reads.update(button["sheet"] for button in procedure_buttons if button["sheet"] in sheets)
            procedures.append({
                "id": f"{Path(module_name).stem.casefold()}::{name.casefold()}",
                "module": module_name,
                "name": name,
                "kind": match.group(1).casefold(),
                "code": body,
                "buttons": procedure_buttons,
                "formula_cells": [],
                "sheets_read": sorted(reads),
                "sheets_written": sorted(writes),
                "unresolved_dynamic_references": list(dict.fromkeys(unresolved))[:20],
            })
    lookup = {row["name"].casefold(): row for row in procedures if row["kind"] != "declarations"}
    for row in procedures:
        row["calls"] = sorted({
            called["id"]
            for name, called in lookup.items()
            if called is not row and re.search(rf"\b{re.escape(name)}\b", row["code"], re.IGNORECASE)
        })
    by_id = {row["id"]: row for row in procedures}
    for _ in procedures:
        changed = False
        for row in procedures:
            for called_id in row["calls"]:
                called = by_id[called_id]
                for key in ("sheets_read", "sheets_written", "unresolved_dynamic_references"):
                    merged = list(dict.fromkeys([*row[key], *called[key]]))
                    if merged != row[key]:
                        row[key] = merged
                        changed = True
        if not changed:
            break
    return procedures


def build_manifest(path: Path) -> dict[str, Any] | None:
    """Return an XLSM-only selection manifest, or None for pass-through files."""
    path = Path(path)
    if not has_vba(path):
        return None
    workbook = load_workbook(path, data_only=False, keep_vba=True, keep_links=False)
    try:
        names = {sheet.title for sheet in workbook.worksheets}
        dependencies = {name: set() for name in names}
        chart_counts = {sheet.title: len(sheet._charts) for sheet in workbook.worksheets}
        image_counts = {sheet.title: len(sheet._images) for sheet in workbook.worksheets}
        charts: dict[str, list[dict[str, Any]]] = {sheet.title: [] for sheet in workbook.worksheets}
        defined = {item.name: _sheet_refs(str(item.attr_text or ""), names) for item in workbook.defined_names.values()}
        defined_lookup = {name.casefold(): refs for name, refs in defined.items()}
        defined_re = re.compile(
            r"(?<![\w.])(" + "|".join(map(re.escape, sorted(defined, key=len, reverse=True))) + r")(?![\w.])",
            re.IGNORECASE,
        ) if defined else None

        def refs(text: str) -> set[str]:
            found = _sheet_refs(text, names)
            if defined_re:
                for match in defined_re.finditer(text):
                    found.update(defined_lookup[match.group(1).casefold()])
            return found

        for sheet in workbook.worksheets:
            for cell in sheet._cells.values():
                formula = getattr(cell.value, "text", cell.value)
                if cell.data_type != "f" or not isinstance(formula, str):
                    continue
                dependencies[sheet.title].update(refs(formula) - {sheet.title})
            for chart in sheet._charts:
                root = chart.to_tree()
                references = list(dict.fromkeys(node.text for node in root.iter() if _local(node.tag) == "f" and node.text))
                dependencies[sheet.title].update(refs(ElementTree.tostring(root, encoding="unicode")) - {sheet.title})
                charts[sheet.title].append({"type": type(chart).__name__.removesuffix("Chart") or "Chart", "references": references})

        buttons = _buttons(path, workbook)
        procedures = _procedures(_modules(path), names, buttons)
        procedure_lookup = {row["name"].casefold(): row for row in procedures if row["kind"] != "declarations"}
        if procedure_lookup:
            called_re = re.compile(r"(?<![\w.])(" + "|".join(map(re.escape, sorted(procedure_lookup, key=len, reverse=True))) + r")\s*\(", re.IGNORECASE)
            for sheet in workbook.worksheets:
                for cell in sheet._cells.values():
                    formula = getattr(cell.value, "text", cell.value)
                    if cell.data_type == "f" and isinstance(formula, str):
                        for match in called_re.finditer(formula):
                            procedure_lookup[match.group(1).casefold()]["formula_cells"].append({"sheet": sheet.title, "cell": cell.coordinate})
        for procedure in procedures:
            for target in procedure["sheets_written"]:
                dependencies[target].update(set(procedure["sheets_read"]) - {target})

        consumers = {name: set() for name in names}
        for consumer, sources in dependencies.items():
            for source in sources:
                consumers[source].add(consumer)
        finals = {
            sheet.title
            for sheet in workbook.worksheets
            if sheet.sheet_state == "visible"
            and (chart_counts[sheet.title] or dependencies[sheet.title] or not consumers[sheet.title])
        }
        if not finals:
            finals = {sheet.title for sheet in workbook.worksheets if sheet.sheet_state == "visible"}

        def lineage(name: str) -> list[str]:
            seen: set[str] = set()
            pending = list(dependencies[name])
            while pending:
                source = pending.pop()
                if source in seen or source == name:
                    continue
                seen.add(source)
                pending.extend(dependencies.get(source, ()))
            return sorted(seen)

        for procedure in procedures:
            affected: set[str] = set()
            pending = list(procedure["sheets_written"])
            seen: set[str] = set()
            while pending:
                sheet = pending.pop()
                if sheet in seen:
                    continue
                seen.add(sheet)
                if sheet in finals:
                    affected.add(sheet)
                pending.extend(consumers.get(sheet, ()))
            procedure["final_outputs_affected"] = sorted(affected)

        return {
            "schema_version": 1,
            "mode": "xlsm-vba-lineage",
            "sheets": [
                {
                    "name": sheet.title,
                    "visibility": sheet.sheet_state,
                    "role": "final_output" if sheet.title in finals else "hidden_intermediate" if sheet.sheet_state != "visible" and consumers[sheet.title] else "source_or_intermediate",
                    "emit": "full" if sheet.title in finals else "lineage",
                    "depends_on": sorted(dependencies[sheet.title]),
                    "lineage": lineage(sheet.title),
                    "chart_count": chart_counts[sheet.title],
                    "charts": charts[sheet.title],
                    "image_count": image_counts[sheet.title],
                }
                for sheet in workbook.worksheets
            ],
            "procedures": procedures,
        }
    finally:
        workbook.close()


def render_consolidated_vba_code(manifest: dict[str, Any] | None) -> str:
    """Consolidated generic-mode VBA-code page.

    Every module and procedure appears once under fenced VB code blocks with
    the static-analysis warning preserved. The page intentionally omits the
    llm-wiki markers (per-procedure ``## シート:`` headings, ``vba://`` links,
    button/formula references, ``<!-- vba-id: -->`` comment markers, and
    sheet/final-output sections) so the generic route stays a plain, whole
    document.
    """
    if not manifest:
        return ""
    lines = [
        "## VBA ソースコード",
        "",
        "> VBAソースコードは静的に抽出しており、マクロは実行していません。",
    ]
    for procedure in manifest.get("procedures") or []:
        module_path = Path(str(procedure.get("module", "")))
        heading = (
            f"### {module_path.name} :: {procedure.get('name', '')}"
            f"（{procedure.get('kind', '')}）"
        )
        lines.extend(["", heading, "", "```vb", str(procedure.get("code", "")), "```"])
    return "\n".join(lines).strip() + "\n"


def render_consolidated_vba_final_outputs(manifest: dict[str, Any] | None) -> str:
    """Consolidated generic-mode VBA-final-output page.

    One compact mapping from every procedure to its final worksheet outputs,
    produced from the same ``final_outputs_affected`` list already carried by
    the manifest. When a procedure cannot statically resolve a sheet
    reference, the unresolved dynamic reference lines are preserved so an
    auditor can see what the mapping did not include.
    """
    if not manifest:
        return ""
    procedures = [
        row
        for row in (manifest.get("procedures") or [])
        if isinstance(row, dict) and row.get("kind") != "declarations"
    ]
    lines = [
        "## VBA 最終出力マッピング",
        "",
        "> VBA は静的解析のみを行い、マクロは実行していません。以下は各プロシージャが"
        "最終的に影響するワークシート出力の対応表です。",
    ]
    for procedure in procedures:
        title = (
            f"{Path(str(procedure.get('module', ''))).name}"
            f" :: {procedure.get('name', '')}"
        )
        lines.extend(["", f"### {title}", ""])
        affected = procedure.get("final_outputs_affected") or []
        lines.extend([f"- `{item}`" for item in affected] or ["- 該当なし"])
        dynamic = procedure.get("unresolved_dynamic_references") or []
        if dynamic:
            lines.extend(["", "静的に解決できない動的参照:", "", *[f"- `{line}`" for line in dynamic]])
    return "\n".join(lines).strip() + "\n"


__all__ = [
    "build_manifest",
    "has_vba",
    "render_consolidated_vba_code",
    "render_consolidated_vba_final_outputs",
]
