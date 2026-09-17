"""Excel parsing with compact HTML tables, formulas, and embedded images."""

from __future__ import annotations

import io
import json
import logging
import os
import posixpath
import re
import subprocess
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from html import escape
from pathlib import Path
from typing import Any
from urllib.parse import quote
from xml.etree import ElementTree

from bs4 import BeautifulSoup, Tag
from oletools.olevba import VBA_Parser
from openpyxl import load_workbook
from openpyxl.utils import coordinate_to_tuple, get_column_letter
from openpyxl.utils.units import (
    DEFAULT_COLUMN_WIDTH,
    DEFAULT_ROW_HEIGHT,
    EMU_to_pixels,
    points_to_pixels,
)
from openpyxl.worksheet.formula import ArrayFormula
from xlsx2html.core import render_table, worksheet_to_data

from client.llm import LLMClient
from formats.base import BaseParser, ExtractedDocument, ParseOptions, ParseProfile
from formats.xlsm_lineage import (
    build_manifest,
    has_vba,
    render_consolidated_vba_code,
    render_consolidated_vba_final_outputs,
)
from utils.markdown_images import embed_markdown_images
from utils.vector_images import convert_document_vector_images, find_libreoffice_command
from workers import Workers

logger = logging.getLogger("doc-parser.xlsx")

_CONTENT_TYPES = "[Content_Types].xml"
_WORKBOOK_XML = "xl/workbook.xml"
_RELATIONSHIP_ID = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
)
_EMBED_ID = (
    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
)
_VECTOR_SUFFIXES = frozenset({".emf", ".wmf", ".svg"})
_PART_MAX_ROWS = 100
_PART_MAX_COLS = 100
_PART_MAX_CHARS = 100_000
_VBA_PROCEDURE_RE = re.compile(
    r"^\s*(?:(?:Public|Private|Friend|Static)\s+)?"
    r"(Sub|Function|Property\s+(?:Get|Let|Set))\s+([^\s(]+)",
    re.IGNORECASE | re.MULTILINE,
)


class OpenpyxlError(RuntimeError):
    """The XLSX renderer could not produce a usable document."""


_SHEET_BLOCK_RE = re.compile(r"(?=^## (?:Sheet|シート): )", re.MULTILINE)


def _split_sheet_blocks(markdown: str) -> list[str]:
    """Split fully-rendered workbook Markdown into per-unit page strings.

    Used only on the llm-wiki path, where the final page content must match the
    image-unit-embedded document. The document title and any preamble before the
    first ``## シート:`` heading remain only in ``markdown``.
    """
    blocks = [
        block.strip()
        for block in _SHEET_BLOCK_RE.split(markdown)
        if block.lstrip().startswith(("## シート: ", "## Sheet: "))
    ]
    return blocks


class LibreOfficeError(RuntimeError):
    """LibreOffice could not recalculate the workbook."""


def recalculate_with_libreoffice(
    xlsx_path: str,
    output_dir: str,
    command: list[str],
) -> str:
    """Open, calculate, and save an XLSX copy through headless LibreOffice."""
    source = Path(xlsx_path).resolve()
    destination = Path(output_dir).resolve()
    destination.mkdir(parents=True, exist_ok=True)
    profile = destination / "libreoffice-profile"
    profile.mkdir(parents=True, exist_ok=True)

    args = [
        *command,
        "--headless",
        "--nologo",
        "--nodefault",
        "--nolockcheck",
        "--nofirststartwizard",
        f"-env:UserInstallation={profile.as_uri()}",
        "--convert-to",
        "xlsx",
        "--outdir",
        str(destination),
        str(source),
    ]
    timeout_s = float(os.getenv("LIBREOFFICE_TIMEOUT_SECONDS", "300"))
    try:
        completed = subprocess.run(
            args,
            capture_output=True,
            text=True,
            timeout=timeout_s,
            check=False,
        )
    except FileNotFoundError as exc:
        raise LibreOfficeError(
            f"LibreOffice executable was not found: {command[0]}"
        ) from exc
    except subprocess.TimeoutExpired as exc:
        raise LibreOfficeError(
            f"LibreOffice recalculation timed out after {timeout_s:g} seconds"
        ) from exc

    recalculated = destination / f"{source.stem}.xlsx"
    if completed.returncode or not recalculated.is_file():
        details = (completed.stderr or completed.stdout or "no output").strip()[-4000:]
        raise LibreOfficeError(
            f"LibreOffice recalculation failed with status "
            f"{completed.returncode}: {details}"
        )
    return str(recalculated)


def _html_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date, time)):
        value = value.isoformat()
    text = escape(str(value), quote=False)
    return text.replace("\r\n", "<br>").replace("\n", "<br>")


def _render_cell(cell, cached_value: Any) -> str:
    if cell.data_type != "f":
        return _html_text(cell.value)

    formula = _html_text(
        cell.value.text if isinstance(cell.value, ArrayFormula) else cell.value
    )
    resolved = (
        _html_text(cached_value)
        if cached_value is not None
        else "[not stored in workbook]"
    )
    return f"数式: <code>{formula}</code><br>キャッシュ値: {resolved}"


def _worksheet_cells(worksheet) -> dict[tuple[int, int], Any]:
    """Return stored, non-empty cells without expanding a sparse used range."""
    return {
        (cell.row, cell.column): cell
        for cell in worksheet._cells.values()
        if cell.value is not None
    }


def _render_worksheet(worksheet, cached_worksheet) -> tuple[str, bool]:
    """Render one worksheet through xlsx2html and remove presentation markup."""
    cells = _worksheet_cells(worksheet)
    if not cells:
        return "_空のシート_", False

    cached_cells = _worksheet_cells(cached_worksheet)
    coordinates = sorted(cells)
    # xlsx2html preserves worksheet layout by walking from A1 through the used
    # range. Avoid expanding pathological sparse ranges into millions of tags.
    area = worksheet.max_row * worksheet.max_column
    max_table_cells = int(os.getenv("XLSX_MAX_TABLE_CELLS", "50000"))
    has_formulas = any(cell.data_type == "f" for cell in cells.values())

    def content(row: int, column: int) -> str:
        cell = cells.get((row, column))
        if cell is None:
            return ""
        cached = cached_cells.get((row, column))
        return _render_cell(cell, cached.value if cached is not None else None)

    if area > max_table_cells:
        lines = [
            "_矩形範囲が非常に大きいため、スパースセル表示を使用しました。_",
            "",
            "<table>",
            "<tr><th>Cell</th><th>Content</th></tr>",
        ]
        for row, column in coordinates:
            coordinate = f"{get_column_letter(column)}{row}"
            lines.append(
                f"<tr><td>{coordinate}</td><td>{content(row, column)}</td></tr>"
            )
        lines.append("</table>")
        return "\n".join(lines), has_formulas

    # Images are emitted separately as document image units. Prevent the
    # renderer from embedding duplicate data URLs inside cells.
    array_formulas = [
        (cell, cell.value)
        for cell in cells.values()
        if isinstance(cell.value, ArrayFormula)
    ]
    for cell, formula in array_formulas:
        cell.value = formula.text
    images = cached_worksheet._images
    cached_worksheet._images = []
    try:
        data = worksheet_to_data(cached_worksheet, fs=worksheet)
    finally:
        cached_worksheet._images = images
        for cell, formula in array_formulas:
            cell.value = formula

    # xlsx2html supplies formatted cached values and handles merged-cell
    # geometry. Keep both the expression and the recalculated value for formula
    # cells so a LibreOffice pass remains visible and auditable.
    for row in data["rows"]:
        for rendered_cell in row:
            formatted = rendered_cell["formatted_value"]
            if formatted == "&nbsp;":
                rendered_cell["formatted_value"] = ""
            elif isinstance(formatted, str):
                rendered_cell["formatted_value"] = formatted.replace(
                    "\r\n", "<br>"
                ).replace("\n", "<br>")

            formula_cell = cells.get(
                (rendered_cell["row"], rendered_cell["column"])
            )
            if formula_cell is None or formula_cell.data_type != "f":
                continue
            cached_cell = cached_cells.get(
                (rendered_cell["row"], rendered_cell["column"])
            )
            rendered_cell["formatted_value"] = _render_cell(
                formula_cell,
                cached_cell.value if cached_cell is not None else None,
            )

    markup = render_table(data, lambda *_: None, lambda *_: None)
    return _simplify_html_table(markup), has_formulas


def _bands(values: list[int], maximum_span: int) -> list[tuple[int, int]]:
    if not values:
        return [(1, 1)]
    output: list[tuple[int, int]] = []
    start = previous = values[0]
    for value in values[1:]:
        if value - start >= maximum_span:
            output.append((start, previous))
            start = value
        previous = value
    output.append((start, previous))
    return output


def _render_worksheet_parts(
    worksheet,
    cached_worksheet,
    extra_positions: list[tuple[int, int]],
) -> tuple[list[tuple[str, tuple[int, int, int, int]]], bool]:
    cells = _worksheet_cells(worksheet)
    positions = [*cells, *extra_positions]
    if not positions:
        return [("_空のシート_", (1, 1, 1, 1))], False

    rows = sorted({row for row, _ in positions})
    cols = sorted({col for _, col in positions})
    rendered = {
        position: _render_cell(
            cell,
            cached_worksheet.cell(*position).value,
        )
        for position, cell in cells.items()
    }
    size = sum(len(value) + 40 for value in rendered.values())
    bounds = (rows[0], rows[-1], cols[0], cols[-1])
    has_formulas = any(cell.data_type == "f" for cell in cells.values())
    if (
        rows[-1] - rows[0] < _PART_MAX_ROWS
        and cols[-1] - cols[0] < _PART_MAX_COLS
        and size <= _PART_MAX_CHARS
    ):
        table, _ = _render_worksheet(worksheet, cached_worksheet)
        if len(table) <= _PART_MAX_CHARS:
            return [(table, bounds)], has_formulas

    parts: list[tuple[str, tuple[int, int, int, int]]] = []
    for row_start, row_end in _bands(rows, _PART_MAX_ROWS):
        for col_start, col_end in _bands(cols, _PART_MAX_COLS):
            selected = [
                (position, value)
                for position, value in rendered.items()
                if row_start <= position[0] <= row_end
                and col_start <= position[1] <= col_end
            ]
            related = [
                position
                for position in extra_positions
                if row_start <= position[0] <= row_end
                and col_start <= position[1] <= col_end
            ]
            if not selected and not related:
                continue
            chunk: list[tuple[tuple[int, int], str]] = []
            chunk_size = 90

            def flush() -> None:
                nonlocal chunk, chunk_size
                if not chunk:
                    return
                lines = [
                    "<table>",
                    "<tr><th>Row</th><th>Cell</th><th>Content</th></tr>",
                ]
                for number, ((row, col), value) in enumerate(chunk, 1):
                    coordinate = f"{get_column_letter(col)}{row}"
                    lines.append(
                        f"<tr><td>{number}</td><td>{coordinate}</td><td>{value}</td></tr>"
                    )
                lines.append("</table>")
                part_rows = [position[0] for position, _ in chunk] or [row_start]
                part_cols = [position[1] for position, _ in chunk] or [col_start]
                parts.append(
                    (
                        "\n".join(lines),
                        (min(part_rows), max(part_rows), min(part_cols), max(part_cols)),
                    )
                )
                chunk = []
                chunk_size = 90

            for position, value in sorted(selected):
                row_size = len(value) + 70
                if chunk and chunk_size + row_size > _PART_MAX_CHARS:
                    flush()
                chunk.append((position, value))
                chunk_size += row_size
            flush()
            if not selected and related:
                parts.append(("_セル値のないオブジェクト範囲_", (row_start, row_end, col_start, col_end)))
    return parts, has_formulas


def _simplify_html_table(markup: str) -> str:
    """Reduce xlsx2html output to semantic table markup without CSS noise."""
    soup = BeautifulSoup(markup, "html.parser")
    table = soup.find("table")
    if not isinstance(table, Tag):
        raise OpenpyxlError("xlsx2html completed without producing a table")

    for tag in list(table.find_all(True)):
        if tag.parent is None:
            continue
        if tag.name in {"col", "colgroup", "img"}:
            tag.decompose()
            continue
        if tag.name not in {"tr", "td", "th", "a", "code", "br"}:
            tag.unwrap()
            continue

        if tag.name in {"td", "th"}:
            tag.attrs = {
                name: value
                for name, value in tag.attrs.items()
                if name in {"rowspan", "colspan"}
                and str(value).isdigit()
                and int(value) > 1
            }
        elif tag.name == "a":
            tag.attrs = {"href": tag["href"]} if tag.get("href") else {}
        else:
            tag.attrs = {}
    table.attrs = {}
    return table.decode(formatter="minimal")


def _column_width_pixels(worksheet, column: int) -> int:
    dimension = worksheet.column_dimensions.get(get_column_letter(column))
    width = (
        dimension.width
        if dimension is not None and dimension.width is not None
        else worksheet.sheet_format.defaultColWidth or DEFAULT_COLUMN_WIDTH
    )
    return max(1, int(width * 7 + 5))


def _row_height_pixels(worksheet, row: int) -> int:
    dimension = worksheet.row_dimensions.get(row)
    height = (
        dimension.height
        if dimension is not None and dimension.height is not None
        else worksheet.sheet_format.defaultRowHeight or DEFAULT_ROW_HEIGHT
    )
    return max(1, points_to_pixels(height))


def _covered_cell_range(
    worksheet,
    start_row: int,
    start_column: int,
    width: int,
    height: int,
    column_offset: int = 0,
    row_offset: int = 0,
) -> str:
    """Return the worksheet cells touched by a one-cell image anchor."""
    end_column = start_column
    remaining_width = max(1, column_offset + width)
    while (
        remaining_width > _column_width_pixels(worksheet, end_column)
        and end_column < 16_384
    ):
        remaining_width -= _column_width_pixels(worksheet, end_column)
        end_column += 1

    end_row = start_row
    remaining_height = max(1, row_offset + height)
    while (
        remaining_height > _row_height_pixels(worksheet, end_row)
        and end_row < 1_048_576
    ):
        remaining_height -= _row_height_pixels(worksheet, end_row)
        end_row += 1

    start = f"{get_column_letter(start_column)}{start_row}"
    end = f"{get_column_letter(end_column)}{end_row}"
    return start if start == end else f"{start}:{end}"


def _image_cell_range(worksheet, image) -> str:
    anchor = image.anchor
    if isinstance(anchor, str):
        row, column = coordinate_to_tuple(anchor)
        return _covered_cell_range(
            worksheet,
            row,
            column,
            round(image.width),
            round(image.height),
        )

    marker = getattr(anchor, "_from", None)
    if marker is None:
        return "不明"
    start_row = marker.row + 1
    start_column = marker.col + 1

    end_marker = getattr(anchor, "to", None)
    if end_marker is not None:
        start = f"{get_column_letter(start_column)}{start_row}"
        end_column = max(start_column, end_marker.col + 1)
        end_row = max(start_row, end_marker.row + 1)
        end = f"{get_column_letter(end_column)}{end_row}"
        return start if start == end else f"{start}:{end}"

    extent = getattr(anchor, "ext", None)
    width = EMU_to_pixels(extent.cx) if extent is not None else round(image.width)
    height = EMU_to_pixels(extent.cy) if extent is not None else round(image.height)
    return _covered_cell_range(
        worksheet,
        start_row,
        start_column,
        width,
        height,
        EMU_to_pixels(marker.colOff),
        EMU_to_pixels(marker.rowOff),
    )


def _extract_images(workbook, output_dir: Path) -> dict[str, list[str]]:
    media_dir = output_dir / "media"
    extracted: dict[str, list[str]] = {}
    image_number = 0

    for worksheet in workbook.worksheets:
        references: list[str] = []
        for image in worksheet._images:
            source_format = (image.format or "png").lower()
            if f".{source_format}" in {".emf", ".wmf", ".svg"}:
                # Preserve these directly from the XLSX package below; Pillow
                # and OpenPyXL cannot reliably decode or re-encode them.
                continue
            image_number += 1
            extension = re.sub(r"[^a-z0-9]", "", source_format)
            extension = extension or "png"
            if extension not in {"gif", "jpeg", "png"}:
                # Image._data() transcodes every other Pillow format to PNG.
                extension = "png"
            filename = f"image-{image_number}.{extension}"
            try:
                payload = image._data()
                media_dir.mkdir(parents=True, exist_ok=True)
                (media_dir / filename).write_bytes(payload)
            except Exception as exc:  # noqa: BLE001
                logger.warning(
                    "could not extract image %s from sheet %s: %s",
                    image_number,
                    worksheet.title,
                    exc,
                )
                continue

            cell_range = _image_cell_range(worksheet, image)
            alt = f"{worksheet.title} シートのセル {cell_range} を覆う画像"
            references.append(
                f"**画像位置:** {worksheet.title} シート、"
                f"セル {cell_range}\n\n"
                f"![{alt}](media/{filename})"
            )
        if references:
            extracted[worksheet.title] = references
    return extracted


def _local_name(tag: str) -> str:
    return tag.rsplit("}", 1)[-1]


def _relationship_targets(
    archive: zipfile.ZipFile,
    part_name: str,
) -> dict[str, str]:
    part_dir, filename = posixpath.split(part_name)
    relationships_name = posixpath.join(part_dir, "_rels", f"{filename}.rels")
    if relationships_name not in archive.namelist():
        return {}

    relationships = ElementTree.fromstring(archive.read(relationships_name))
    targets: dict[str, str] = {}
    for relationship in relationships:
        if relationship.attrib.get("TargetMode") == "External":
            continue
        target = relationship.attrib.get("Target")
        relationship_id = relationship.attrib.get("Id")
        if not target or not relationship_id:
            continue
        if target.startswith("/"):
            resolved = target.lstrip("/")
        else:
            resolved = posixpath.normpath(posixpath.join(part_dir, target))
        targets[relationship_id] = resolved
    return targets


def _extract_vba_modules(xlsx_path: Path) -> list[tuple[str, str]]:
    with zipfile.ZipFile(xlsx_path) as archive:
        if "xl/vbaProject.bin" not in archive.namelist():
            return []

    parser = VBA_Parser(str(xlsx_path))
    try:
        return [
            (Path(module_name).name, str(code).replace("\r\n", "\n"))
            for _, _, module_name, code in parser.extract_macros()
        ]
    except Exception as exc:
        raise OpenpyxlError(f"could not extract VBA source: {exc}") from exc
    finally:
        parser.close()


def _vba_pages(modules: list[tuple[str, str]]) -> list[dict[str, str]]:
    pages: list[dict[str, str]] = []
    for module_name, code in modules:
        matches = list(_VBA_PROCEDURE_RE.finditer(code))
        declarations = "\n".join(
            line
            for line in code[: matches[0].start() if matches else len(code)].splitlines()
            if line.strip() and not line.lstrip().startswith("Attribute VB_")
        )
        module = Path(module_name).stem
        if declarations:
            pages.append(
                {
                    "id": f"{module.casefold()}::declarations",
                    "title": f"マクロ-{module}-宣言",
                    "module": module_name,
                    "name": "宣言",
                    "kind": "declarations",
                    "code": declarations,
                }
            )
        for index, match in enumerate(matches):
            name = match.group(2)
            pages.append(
                {
                    "id": f"{module.casefold()}::{name.casefold()}",
                    "title": f"マクロ-{module}-{name}",
                    "module": module_name,
                    "name": name,
                    "kind": match.group(1).casefold(),
                    "code": code[
                        match.start() : matches[index + 1].start()
                        if index + 1 < len(matches)
                        else len(code)
                    ].rstrip(),
                }
            )
    return pages


def _vba_lookup(pages: list[dict[str, str]]) -> dict[str, dict[str, str]]:
    output: dict[str, dict[str, str]] = {}
    for page in pages:
        if page["kind"] != "declarations":
            output.setdefault(page["name"].casefold(), page)
            output[f'{Path(page["module"]).stem.casefold()}.{page["name"].casefold()}'] = page
    return output


def _render_vba_page(
    page: dict[str, str],
    usages: list[str],
    details: dict[str, Any] | None = None,
) -> str:
    output = [
        f'<!-- vba-id: {quote(page["id"], safe="")} -->',
        "> VBAソースコードは静的に抽出しており、マクロは実行していません。",
        "",
        f'- ソースモジュール: `{page["module"]}`',
        f'- 種別: `{page["kind"]}`',
    ]
    if usages:
        output.extend(["", "## ワークブック内の参照元", "", *usages])
    if details:
        manifest_buttons = [
            f"- `{button.get('sheet', '不明')}!{button.get('cells', '不明')}` のボタン"
            for button in details.get("buttons") or []
        ]
        output.extend(["", "## ボタン・セル", "", *(manifest_buttons or ["- 該当なし"])])
        for heading, key in (
            ("参照シート", "sheets_read"),
            ("更新シート", "sheets_written"),
            ("影響する最終出力", "final_outputs_affected"),
        ):
            values = details.get(key) or []
            output.extend(["", f"## {heading}", "", *([f"- `{value}`" for value in values] or ["- 該当なし"])])
        dynamic = details.get("unresolved_dynamic_references") or []
        if dynamic:
            output.extend(
                [
                    "",
                    "## 静的に解決できない動的参照",
                    "",
                    "> 静的に参照先を確定できないため、参照元シートを推測で展開していません。",
                    "",
                    *[f"- `{line}`" for line in dynamic],
                ]
            )
    output.extend(["", "## コード", "", "```vb", page["code"], "```"])
    return "\n".join(output)


def _chart_summary(worksheet) -> list[str]:
    output: list[str] = []
    for number, chart in enumerate(worksheet._charts, 1):
        root = chart.to_tree()
        formulas = list(dict.fromkeys(node.text for node in root.iter() if _local_name(node.tag) == "f" and node.text))
        chart_type = type(chart).__name__.removesuffix("Chart") or "Chart"
        source = ", ".join(f"`{value}`" for value in formulas) or "参照範囲不明"
        output.append(f"- グラフ {number}（{chart_type}）: {source}")
    return output


def _selection(manifest_json: str | None, workbook) -> tuple[set[str] | None, dict[str, dict[str, Any]]]:
    if not manifest_json:
        return None, {}
    try:
        manifest = json.loads(manifest_json)
        if manifest.get("mode") != "xlsm-vba-lineage":
            return None, {}
        sheets = {row["name"]: row for row in manifest.get("sheets", []) if isinstance(row, dict) and row.get("name")}
        selected = {name for name, row in sheets.items() if row.get("emit") == "full"}
        if not selected or not selected.issubset(set(workbook.sheetnames)):
            raise OpenpyxlError("invalid XLSM sheet-selection manifest")
        return selected, sheets
    except (KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        raise OpenpyxlError(f"invalid XLSM sheet-selection manifest: {exc}") from exc


def _sheet_context(
    name: str,
    sheet_manifest: dict[str, dict[str, Any]],
    charts: dict[str, list[str]],
) -> list[str]:
    row = sheet_manifest.get(name, {})
    direct = row.get("depends_on") or []
    lineage = row.get("lineage") or []
    related = [source for source in lineage if source in sheet_manifest]
    output = ["<!-- sheet-context:start -->"]
    if related:
        output.extend(
            [
                "### シートの系譜",
                "",
                f"- 最終出力: `{name}`",
                f"- 直接入力: {', '.join(f'`{value}`' for value in direct)}",
                f"- 上流シート: {', '.join(f'`{value}`' for value in related)}",
            ]
        )
    attached_charts = [*charts.get(name, [])]
    if attached_charts:
        output.extend(["", "### グラフ", "", *attached_charts])
    output.append("<!-- sheet-context:end -->")
    return output if len(output) > 2 else []


def _vba_control_locations(
    archive: zipfile.ZipFile,
    workbook,
) -> dict[str, list[tuple[str, str]]]:
    workbook_rels = _relationship_targets(archive, _WORKBOOK_XML)
    workbook_xml = ElementTree.fromstring(archive.read(_WORKBOOK_XML))
    worksheets = {worksheet.title: worksheet for worksheet in workbook.worksheets}
    locations: dict[str, list[tuple[str, str]]] = {}

    for sheet in workbook_xml.iter():
        if _local_name(sheet.tag) != "sheet":
            continue
        sheet_name = sheet.attrib.get("name", "不明なシート")
        worksheet = worksheets.get(sheet_name)
        sheet_part = workbook_rels.get(sheet.attrib.get(_RELATIONSHIP_ID, ""))
        if (
            worksheet is None
            or sheet_part is None
            or sheet_part not in archive.namelist()
        ):
            continue
        sheet_rels = _relationship_targets(archive, sheet_part)
        sheet_xml = ElementTree.fromstring(archive.read(sheet_part))
        for drawing in sheet_xml.iter():
            if _local_name(drawing.tag) not in {"drawing", "legacyDrawing"}:
                continue
            drawing_part = sheet_rels.get(drawing.attrib.get(_RELATIONSHIP_ID, ""))
            if drawing_part is None or drawing_part not in archive.namelist():
                continue
            drawing_xml = ElementTree.fromstring(archive.read(drawing_part))
            if _local_name(drawing.tag) == "drawing":
                for anchor in drawing_xml.iter():
                    if _local_name(anchor.tag) not in {
                        "oneCellAnchor",
                        "twoCellAnchor",
                        "absoluteAnchor",
                    }:
                        continue
                    macro = next(
                        (
                            value
                            for node in anchor.iter()
                            for key, value in node.attrib.items()
                            if _local_name(key) == "macro" and value
                        ),
                        "",
                    )
                    if macro:
                        locations.setdefault(sheet_name, []).append(
                            (
                                _drawing_anchor_range(anchor, worksheet),
                                macro.rsplit("!", 1)[-1],
                            )
                        )
                continue

            for client in drawing_xml.iter():
                if _local_name(client.tag) != "ClientData":
                    continue
                values = {
                    _local_name(child.tag): (child.text or "").strip()
                    for child in client
                }
                macro = values.get("FmlaMacro", "").rsplit("!", 1)[-1]
                try:
                    anchor = [int(value.strip()) for value in values["Anchor"].split(",")]
                    start = f"{get_column_letter(anchor[0] + 1)}{anchor[2] + 1}"
                    end = f"{get_column_letter(anchor[4] + 1)}{anchor[6] + 1}"
                except (KeyError, ValueError, IndexError):
                    continue
                if macro:
                    locations.setdefault(sheet_name, []).append(
                        (start if start == end else f"{start}:{end}", macro)
                    )
    return locations


def _render_vba_references(
    worksheet,
    controls: list[tuple[str, str]],
    procedures: dict[str, dict[str, str]],
    usages: dict[str, list[str]],
) -> list[tuple[tuple[int, int], str]]:
    references: list[tuple[tuple[int, int], str]] = []
    for cell_range, macro in controls:
        procedure = procedures.get(macro.casefold()) or procedures.get(
            macro.rsplit(".", 1)[-1].casefold()
        )
        if procedure:
            position = coordinate_to_tuple(cell_range.split(":", 1)[0])
            target = f'vba://{quote(procedure["id"], safe="")}'
            references.append(
                (
                    position,
                    f"- ボタン範囲 **{cell_range}**: [{procedure['name']}]({target})",
                )
            )
            usages.setdefault(procedure["id"], []).append(
                f"- Button on `{worksheet.title}!{cell_range}`"
            )

    functions = {
        name: value
        for name, value in procedures.items()
        if "." not in name and value["kind"] == "function"
    }
    if functions:
        pattern = re.compile(
            r"(?<![\w.])(" + "|".join(map(re.escape, functions)) + r")\s*\(",
            re.IGNORECASE,
        )
        for cell in worksheet._cells.values():
            formula = getattr(cell.value, "text", cell.value)
            if cell.data_type != "f" or not isinstance(formula, str):
                continue
            for match in pattern.finditer(formula):
                procedure = functions[match.group(1).casefold()]
                target = f'vba://{quote(procedure["id"], safe="")}'
                references.append(
                    (
                        (cell.row, cell.column),
                        f"- 数式セル **{cell.coordinate}**: "
                        f"[{procedure['name']}]({target})",
                    )
                )
                usages.setdefault(procedure["id"], []).append(
                    f"- Formula in `{worksheet.title}!{cell.coordinate}`"
                )
    return list(dict.fromkeys(references))


def _reference_position(text: str) -> tuple[int, int] | None:
    match = re.search(r"(?:セル |\*\*)([A-Z]+\d+)", text)
    return coordinate_to_tuple(match.group(1)) if match else None


def _references_by_part(
    parts: list[tuple[str, tuple[int, int, int, int]]],
    references: list[tuple[tuple[int, int] | None, str]],
) -> list[list[str]]:
    output = [[] for _ in parts]
    for position, text in references:
        index = 0
        if position is not None:
            row, col = position
            index = next(
                (
                    number
                    for number, (_, (r1, r2, c1, c2)) in enumerate(parts)
                    if r1 <= row <= r2 and c1 <= col <= c2
                ),
                0,
            )
        output[index].append(text)
    return output


def _marker_values(node) -> tuple[int, int, int, int] | None:
    values: dict[str, int] = {}
    for child in node:
        if child.text is None:
            continue
        try:
            values[_local_name(child.tag)] = int(child.text)
        except ValueError:
            continue
    if "row" not in values or "col" not in values:
        return None
    return (
        values["row"] + 1,
        values["col"] + 1,
        EMU_to_pixels(values.get("rowOff", 0)),
        EMU_to_pixels(values.get("colOff", 0)),
    )


def _drawing_anchor_range(anchor, worksheet) -> str:
    start_node = next(
        (child for child in anchor if _local_name(child.tag) == "from"),
        None,
    )
    if start_node is None or (start := _marker_values(start_node)) is None:
        return "不明"
    start_row, start_column, row_offset, column_offset = start

    end_node = next(
        (child for child in anchor if _local_name(child.tag) == "to"),
        None,
    )
    if end_node is not None and (end := _marker_values(end_node)) is not None:
        end_row, end_column, _, _ = end
        start_cell = f"{get_column_letter(start_column)}{start_row}"
        end_cell = f"{get_column_letter(end_column)}{end_row}"
        return start_cell if start_cell == end_cell else f"{start_cell}:{end_cell}"

    extent = next(
        (child for child in anchor if _local_name(child.tag) == "ext"),
        None,
    )
    if extent is None:
        return f"{get_column_letter(start_column)}{start_row}"
    return _covered_cell_range(
        worksheet,
        start_row,
        start_column,
        EMU_to_pixels(int(extent.attrib.get("cx", "0"))),
        EMU_to_pixels(int(extent.attrib.get("cy", "0"))),
        column_offset,
        row_offset,
    )


def _vector_image_locations(
    archive: zipfile.ZipFile,
    workbook,
) -> dict[str, list[tuple[str, str]]]:
    """Map vector-media package members to their sheet and covered cells."""
    workbook_rels = _relationship_targets(archive, _WORKBOOK_XML)
    workbook_xml = ElementTree.fromstring(archive.read(_WORKBOOK_XML))
    worksheets = {worksheet.title: worksheet for worksheet in workbook.worksheets}
    locations: dict[str, list[tuple[str, str]]] = {}

    for sheet in workbook_xml.iter():
        if _local_name(sheet.tag) != "sheet":
            continue
        sheet_name = sheet.attrib.get("name", "不明なシート")
        worksheet = worksheets.get(sheet_name)
        sheet_part = workbook_rels.get(sheet.attrib.get(_RELATIONSHIP_ID, ""))
        if (
            worksheet is None
            or sheet_part is None
            or sheet_part not in archive.namelist()
        ):
            continue

        sheet_rels = _relationship_targets(archive, sheet_part)
        sheet_xml = ElementTree.fromstring(archive.read(sheet_part))
        for drawing in sheet_xml.iter():
            if _local_name(drawing.tag) != "drawing":
                continue
            drawing_part = sheet_rels.get(drawing.attrib.get(_RELATIONSHIP_ID, ""))
            if drawing_part is None or drawing_part not in archive.namelist():
                continue

            drawing_rels = _relationship_targets(archive, drawing_part)
            drawing_xml = ElementTree.fromstring(archive.read(drawing_part))
            for anchor in drawing_xml.iter():
                if _local_name(anchor.tag) not in {
                    "oneCellAnchor",
                    "twoCellAnchor",
                    "absoluteAnchor",
                }:
                    continue
                cell_range = _drawing_anchor_range(anchor, worksheet)
                for node in anchor.iter():
                    if _local_name(node.tag) != "blip":
                        continue
                    media_part = drawing_rels.get(node.attrib.get(_EMBED_ID, ""))
                    if (
                        media_part is not None
                        and Path(media_part).suffix.lower() in _VECTOR_SUFFIXES
                    ):
                        locations.setdefault(media_part, []).append(
                            (sheet_name, cell_range)
                        )
    return locations


def _extract_package_vector_images(
    xlsx_path: Path,
    output_dir: Path,
    workbook,
) -> dict[str, list[str]]:
    """Extract dropped vector media with its worksheet cell placement."""
    media_dir = output_dir / "media"
    references: dict[str, list[str]] = {}

    with zipfile.ZipFile(xlsx_path) as archive:
        members = sorted(
            name
            for name in archive.namelist()
            if name.startswith("xl/media/")
            and Path(name).suffix.lower() in _VECTOR_SUFFIXES
        )
        locations = _vector_image_locations(archive, workbook)
        for number, member in enumerate(members, start=1):
            suffix = Path(member).suffix.lower()
            filename = f"vector-{number}{suffix}"
            media_dir.mkdir(parents=True, exist_ok=True)
            (media_dir / filename).write_bytes(archive.read(member))
            placements = locations.get(member) or [("", "不明")]
            for sheet_name, cell_range in placements:
                alt = (
                    f"{sheet_name} シートのセル {cell_range} を覆う"
                    "ベクター画像"
                    if sheet_name
                    else f"位置不明の Excel ベクター画像 {number}"
                )
                location = (
                    f"{sheet_name} シート、セル {cell_range}"
                    if sheet_name
                    else "不明"
                )
                references.setdefault(sheet_name, []).append(
                    f"**画像位置:** {location}\n\n![{alt}](media/{filename})"
                )
    return references


@dataclass(slots=True)
class RenderedWorkbook:
    """Picklable result of rendering a workbook into worksheet Markdown units.

    ``pages`` holds raw, un-embedded unit strings in assembly order. In the
    generic profile each page is one whole worksheet (plus consolidated VBA
    pages for XLSM); llm-wiki leaves ``pages`` empty because its final page
    form must be derived after the image-unit LLM pipeline runs.
    """

    markdown: str
    pages: list[str] = field(default_factory=list)
    markdown_path: Path | None = None
    asset_root: Path | None = None


def _generic_render(
    formulas,
    cached_values,
    images: dict[str, list[str]],
    vector_images: dict[str, list[str]],
    macro_source: Path,
    is_macro: bool,
    title: str,
    has_formulas: bool,
) -> tuple[str, list[str]]:
    """Render one whole unit per worksheet (+ consolidated VBA pages) generically.

    No worksheet splitting, no lineage/chart context, and no ``vba://`` links.
    Static VBA analysis runs on the original upload and produces exactly two
    synthetic pages: one consolidated code page and one final-output page.
    """
    preamble = [f"# {_html_text(title)}"]
    if has_formulas:
        preamble.append(
            "> 数式セルには数式とワークブックのキャッシュ値の両方を記載しています。"
            "利用可能な環境では LibreOffice により抽出前に再計算されます。"
        )
    units: list[str] = []
    for worksheet in formulas.worksheets:
        cached_worksheet = cached_values[worksheet.title]
        table, _ = _render_worksheet(worksheet, cached_worksheet)
        block = [
            f"## シート: {_html_text(worksheet.title)}",
            "",
            table,
        ]
        sheet_images = [*images.get(worksheet.title, []), *vector_images.get(worksheet.title, [])]
        if sheet_images:
            block.extend(["", "### 画像", "", *sheet_images])
        units.append("\n".join(block))

    unplaced = vector_images.get("")
    if unplaced:
        # ``pages`` is one item per worksheet. Keep package media whose sheet
        # anchor could not be resolved inside the first sheet instead of
        # inventing a synthetic worksheet page.
        section = "\n".join(["### 位置不明のベクター画像", "", *unplaced])
        if units:
            units[0] = f"{units[0]}\n\n{section}"

    manifest = build_manifest(macro_source) if is_macro else None
    code_page = render_consolidated_vba_code(manifest)
    final_page = render_consolidated_vba_final_outputs(manifest)
    if code_page:
        units.append(code_page.rstrip("\n"))
    if final_page:
        units.append(final_page.rstrip("\n"))

    markdown = "\n\n".join(["\n".join(preamble), *units])
    return markdown, units


def run_openpyxl(
    xlsx_path: str,
    output_dir: str,
    document_title: str | None = None,
    vba_path: str | None = None,
    manifest_json: str | None = None,
    profile: ParseProfile = ParseProfile.LLM_WIKI,
    is_macro: bool = False,
) -> RenderedWorkbook:
    """Convert workbook cells to compact HTML tables and extract its images.

    ``profile=GENERIC`` renders every worksheet whole (never split), attaches
    each sheet's images to that same unit, and emits two consolidated static
    VBA pages for macro workbooks. ``profile=LLM_WIKI`` preserves the existing
    manifest-driven selection, splitting, lineage, and per-procedure output and
    defers page construction to the parser.
    """
    source = Path(xlsx_path).resolve()
    macro_source = Path(vba_path).resolve() if vba_path else source
    destination = Path(output_dir).resolve()
    destination.mkdir(parents=True, exist_ok=True)
    markdown_path = destination / "document.md"

    try:
        formulas = load_workbook(source, data_only=False, keep_links=False)
        cached_values = load_workbook(source, data_only=True, keep_links=False)
    except Exception as exc:
        raise OpenpyxlError(f"OpenPyXL could not read the workbook: {exc}") from exc

    try:
        images = _extract_images(formulas, destination)
        vector_images = _extract_package_vector_images(source, destination, formulas)
        title = document_title or "Excel ワークブック"

        # The generic profile never consults the manifest or splits worksheets,
        # and it never renders vba:// links or per-procedure pages.
        if profile == ParseProfile.GENERIC:
            sheet_formulas = any(
                cell.data_type == "f"
                for sheet in formulas.worksheets
                for cell in sheet._cells.values()
            )
            markdown, pages = _generic_render(
                formulas,
                cached_values,
                images,
                vector_images,
                macro_source,
                is_macro,
                title,
                sheet_formulas,
            )
            markdown_path.write_text(markdown + "\n", encoding="utf-8")
            return RenderedWorkbook(
                markdown=markdown,
                pages=pages,
                markdown_path=markdown_path,
                asset_root=destination,
            )

        vba_modules = _extract_vba_modules(macro_source)
        vba_pages = _vba_pages(vba_modules)
        vba_procedures = _vba_lookup(vba_pages)
        vba_usages: dict[str, list[str]] = {}
        with zipfile.ZipFile(macro_source) as archive:
            vba_controls = _vba_control_locations(archive, formulas)
        selected, sheet_manifest = _selection(manifest_json, formulas)
        chart_summaries = (
            {sheet.title: _chart_summary(sheet) for sheet in formulas.worksheets}
            if selected is not None
            else {}
        )
        sections = [f"# {_html_text(title)}"]
        has_formulas = False
        for worksheet in formulas.worksheets:
            if selected is not None and worksheet.title not in selected:
                continue
            cached_worksheet = cached_values[worksheet.title]
            vba_references = _render_vba_references(
                worksheet,
                vba_controls.get(worksheet.title, []),
                vba_procedures,
                vba_usages,
            )
            image_references = [
                (position, reference)
                for reference in [
                    *images.get(worksheet.title, []),
                    *vector_images.get(worksheet.title, []),
                ]
                if (position := _reference_position(reference)) is not None
            ]
            parts, sheet_has_formulas = _render_worksheet_parts(
                worksheet,
                cached_worksheet,
                [position for position, _ in [*image_references, *vba_references]],
            )
            has_formulas = has_formulas or sheet_has_formulas
            part_images = _references_by_part(parts, image_references)
            part_vba = _references_by_part(parts, vba_references)
            for index, ((table, bounds), page_images, page_vba) in enumerate(
                zip(parts, part_images, part_vba),
                1,
            ):
                page_title = (
                    worksheet.title
                    if len(parts) == 1
                    else f"{worksheet.title}-part{index}"
                )
                r1, r2, c1, c2 = bounds
                sections.extend(
                    [
                        "",
                        f"## シート: {_html_text(page_title)}",
                        "",
                        f"_元範囲: {get_column_letter(c1)}{r1}:{get_column_letter(c2)}{r2}_",
                        "",
                        table,
                    ]
                )
                if page_images:
                    sections.extend(["", "### 画像", "", *page_images])
                if page_vba:
                    sections.extend(
                        [
                            "",
                            "<!-- vba-references:start -->",
                            "### VBA参照",
                            "",
                            *page_vba,
                            "<!-- vba-references:end -->",
                        ]
                    )
                if selected is not None and index == 1:
                    context = _sheet_context(
                        worksheet.title,
                        sheet_manifest,
                        chart_summaries,
                    )
                    if context:
                        sections.extend(["", *context])

        unplaced_vectors = vector_images.get("") if selected is None else None
        if unplaced_vectors:
            sections.extend(
                ["", "## シート: 位置不明のベクター画像", "", *unplaced_vectors]
            )

        manifest_procedures = {
            row.get("id"): row
            for row in (json.loads(manifest_json).get("procedures", []) if manifest_json else [])
            if isinstance(row, dict) and row.get("id")
        }
        for page in vba_pages:
            sections.extend(
                [
                    "",
                    f'## シート: {_html_text(page["title"])}',
                    "",
                    _render_vba_page(page, vba_usages.get(page["id"], []), manifest_procedures.get(page["id"])),
                ]
            )

        if has_formulas:
            sections[1:1] = [
                "",
                (
                    "> 数式セルには数式とワークブックのキャッシュ値の両方を記載しています。"
                    "利用可能な環境では LibreOffice により抽出前に再計算されます。"
                ),
            ]
        markdown = "\n".join(sections) + "\n"
        markdown_path.write_text(markdown, encoding="utf-8")
    except OpenpyxlError:
        raise
    except Exception as exc:
        raise OpenpyxlError(f"could not render the workbook: {exc}") from exc
    finally:
        formulas.close()
        cached_values.close()

    return RenderedWorkbook(markdown=markdown, pages=[], markdown_path=markdown_path, asset_root=destination)


class XlsxParser(BaseParser):
    name = "xlsx"

    @classmethod
    def detect(cls, data: bytes) -> bool:
        if not data.startswith(b"PK"):
            return False
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as archive:
                names = set(archive.namelist())
        except (OSError, zipfile.BadZipFile):
            return False
        return _CONTENT_TYPES in names and _WORKBOOK_XML in names

    async def _extract(
        self,
        data: bytes,
        image_dir: str,
        options: ParseOptions,
        workers: Workers,
    ) -> ExtractedDocument:
        work_dir = Path(image_dir)
        upload_name = Path(options.filename or "").name or "document.xlsx"
        if Path(upload_name).suffix.lower() not in {".xlsx", ".xlsm"}:
            upload_name = f"{Path(upload_name).stem}.xlsx"
        xlsx_path = work_dir / upload_name
        output_dir = work_dir / "openpyxl-output"
        xlsx_path.write_bytes(data)
        document_title = (
            Path(upload_name).stem
            if Path(upload_name).stem not in {"document", ""}
            else "Excel ワークブック"
        )
        is_macro = has_vba(xlsx_path)
        generic = options.profile == ParseProfile.GENERIC

        # Static lineage for llm-wiki when a manifest is missing must originate
        # from the original upload; the caller-supplied manifest is preserved
        # verbatim so any downstream ordering/translation stays intact.
        manifest = options.manifest
        if not generic and is_macro and manifest is None:
            manifest = build_manifest(xlsx_path)

        workbook_path = xlsx_path
        recalculate_mode = os.getenv("XLSX_RECALCULATE_FORMULAS", "auto").lower()
        skip_recalculation = is_macro or bool(
            manifest and manifest.get("mode") == "xlsm-vba-lineage"
        )
        if not skip_recalculation and recalculate_mode not in {"0", "false", "no", "off"}:
            command = find_libreoffice_command()
            if command is None:
                if recalculate_mode in {"1", "true", "yes", "required"}:
                    raise OpenpyxlError(
                        "formula recalculation was required, but LibreOffice "
                        "was not found"
                    )
            else:
                try:
                    workbook_path = Path(
                        await workers.run_external(
                            recalculate_with_libreoffice,
                            str(xlsx_path),
                            str(work_dir / "recalculated"),
                            command,
                        )
                    )
                except LibreOfficeError as exc:
                    if recalculate_mode in {"1", "true", "yes", "required"}:
                        raise
                    logger.warning(
                        "LibreOffice recalculation failed; using saved values: %s",
                        exc,
                    )

        rendered: RenderedWorkbook = await workers.run_external(
            run_openpyxl,
            str(workbook_path),
            str(output_dir),
            document_title,
            str(xlsx_path),
            json.dumps(manifest, ensure_ascii=False) if manifest else None,
            options.profile,
            is_macro,
        )

        if generic:
            # Convert once across the full document and its duplicated page
            # views. Re-running LibreOffice per page creates different PNG
            # names and can make ``markdown`` disagree with ``pages``.
            boundary = "\n<!-- doc-parser-page-boundary -->\n"
            parts = [rendered.markdown, *rendered.pages]
            if any(boundary in part for part in parts):
                raise OpenpyxlError("reserved page boundary appeared in workbook output")
            converted = await convert_document_vector_images(
                boundary.join(parts), output_dir / "media", workers
            )
            markdown, *pages = converted.split(boundary)
            return ExtractedDocument(
                markdown=markdown,
                pages=pages,
                markdown_path=rendered.markdown_path,
                asset_root=rendered.asset_root,
            )

        # llm-wiki path: preserve the existing image-unit LLM pipeline exactly.
        markdown = rendered.markdown
        markdown_path = rendered.markdown_path
        assert markdown_path is not None
        markdown = await convert_document_vector_images(
            markdown, output_dir / "media", workers
        )

        client = (
            LLMClient(
                base_url=options.llm_base_url,
                api_key=options.llm_api_key,
                model=options.llm_model,
            )
            if options.describe_images
            else None
        )
        try:
            embedded = await embed_markdown_images(
                markdown,
                markdown_path,
                output_dir,
                workers,
                client.describe_image if client is not None else None,
            )
        finally:
            if client is not None:
                await client.close()
        # Derive pages from the fully-embedded final llm-wiki output so their
        # shape matches markdown. The document title stays only in markdown.
        pages = _split_sheet_blocks(embedded)
        return ExtractedDocument(markdown=embedded, pages=pages)
