from __future__ import annotations

import asyncio
import base64
import io
import json
import shutil
import subprocess
import tempfile
import unittest
import warnings
import zipfile
from pathlib import Path
from typing import ClassVar
from unittest.mock import patch

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as SpreadsheetImage
from openpyxl.worksheet.formula import ArrayFormula
from PIL import Image

from formats import detect
from formats.base import ParseOptions
from formats.xlsx import (
    XlsxParser,
    _render_worksheet,
    _render_worksheet_parts,
    _render_vba_page,
    _render_vba_references,
    _vba_lookup,
    _vba_pages,
    recalculate_with_libreoffice,
    run_openpyxl,
)

_PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def make_xlsx() -> bytes:
    stream = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Summary"
    sheet["A1"] = "Revenue"
    sheet["B1"] = 10
    sheet["B2"] = 20
    sheet["B3"] = "=SUM(B1:B2)"
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class FakeWorkers:
    def __init__(self) -> None:
        self.external_functions = []
        self.network_calls = 0

    async def run_external(self, fn, source: str, output_dir: str, *args) -> str:
        self.external_functions.append(fn)
        destination = Path(output_dir)
        destination.mkdir(parents=True, exist_ok=True)

        if fn is recalculate_with_libreoffice:
            recalculated = destination / Path(source).name
            shutil.copyfile(source, recalculated)
            return str(recalculated)

        media = destination / "media"
        media.mkdir()
        (media / "image-1.png").write_bytes(_PNG_BYTES)
        markdown_path = destination / "document.md"
        markdown_path.write_text(
            "# Excel workbook\n\n"
            "<table>\n"
            "<tr><td>Revenue</td><td>10</td></tr>\n"
            "</table>\n\n"
            "**画像位置:** Summary シート、セル D4\n\n"
            "![Summary シートのセル D4 を覆う画像](media/image-1.png)\n",
            encoding="utf-8",
        )
        return str(markdown_path)

    async def run_network(self, fn, *args, **kwargs):
        self.network_calls += 1
        return await fn(*args, **kwargs)


class FakeLLMClient:
    configurations: ClassVar[list[dict]] = []

    def __init__(self, **kwargs) -> None:
        self.configurations.append(kwargs)

    async def describe_image(self, data_url: str, alt_text: str) -> str:
        await asyncio.sleep(0)
        return f"Description of {alt_text}"

    async def close(self) -> None:
        return None


class XlsxParserTests(unittest.IsolatedAsyncioTestCase):
    def test_detects_excel_without_claiming_docx(self) -> None:
        xlsx = make_xlsx()

        self.assertTrue(XlsxParser.detect(xlsx))
        self.assertIs(detect(xlsx), XlsxParser)
        self.assertFalse(XlsxParser.detect(b"PK not really an xlsx"))

    def test_formula_cell_contains_formula_and_cached_result(self) -> None:
        formulas = Workbook()
        formulas.active["A1"] = 2
        formulas.active["A2"] = 3
        formulas.active["A3"] = "=SUM(A1:A2)"
        cached = Workbook()
        cached.active["A1"] = 2
        cached.active["A2"] = 3
        cached.active["A3"] = 5

        html, has_formulas = _render_worksheet(formulas.active, cached.active)

        self.assertTrue(has_formulas)
        self.assertIn("数式: <code>=SUM(A1:A2)</code>", html)
        self.assertIn("キャッシュ値: 5", html)
        formulas.close()
        cached.close()

    def test_array_formula_is_rendered(self) -> None:
        formulas = Workbook()
        formulas.active["A1"] = ArrayFormula("A1:A2", "=SUM(B1:B2)")
        cached = Workbook()
        cached.active["A1"] = 5

        html, has_formulas = _render_worksheet(formulas.active, cached.active)

        self.assertTrue(has_formulas)
        self.assertIn("数式: <code>=SUM(B1:B2)</code>", html)
        self.assertIn("キャッシュ値: 5", html)
        formulas.close()
        cached.close()

    def test_vba_source_and_formula_reference_link_to_vba_page(self) -> None:
        modules = [
            (
                "Module1.bas",
                "Function DoubleIt(value)\nDoubleIt = value * 2\nEnd Function\n"
                "Sub SecondMacro()\nEnd Sub\n",
            )
        ]
        workbook = Workbook()
        workbook.active["B2"] = "=DoubleIt(A2)"
        pages = _vba_pages(modules)
        usages: dict[str, list[str]] = {}

        self.assertEqual(len(pages), 2)
        self.assertNotIn("SecondMacro", pages[0]["code"])
        self.assertNotIn("DoubleIt", pages[1]["code"])

        references = _render_vba_references(
            workbook.active,
            [("D4:F5", "DoubleIt")],
            _vba_lookup(pages),
            usages,
        )
        page = _render_vba_page(pages[0], usages[pages[0]["id"]])

        self.assertIn("```vb\nFunction DoubleIt", page)
        self.assertIn("Button on `Sheet!D4:F5`", page)
        self.assertIn("Formula in `Sheet!B2`", page)
        self.assertEqual(
            references,
            [
                ((4, 4), "- ボタン範囲 **D4:F5**: [DoubleIt](vba://module1%3A%3Adoubleit)"),
                ((2, 2), "- 数式セル **B2**: [DoubleIt](vba://module1%3A%3Adoubleit)"),
            ],
        )
        workbook.close()

    def test_manifest_emits_sheets_before_vba_with_lineage_charts_and_images(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            output = root / "output"
            workbook = Workbook()
            settings = workbook.active
            settings.title = "設定"
            data = workbook.create_sheet("グラフデータ")
            data.sheet_state = "hidden"
            for row in range(1, 244):
                data.cell(row, 1, row)
            image_stream = io.BytesIO()
            Image.new("RGB", (2, 2), "blue").save(image_stream, format="PNG")
            image_stream.seek(0)
            data.add_image(SpreadsheetImage(image_stream), "B2")
            graph = workbook.create_sheet("グラフ")
            chart = BarChart()
            chart.add_data(Reference(data, min_col=1, min_row=1, max_row=243))
            graph.add_chart(chart, "A1")
            workbook.save(source)
            workbook.close()
            manifest = {
                "mode": "xlsm-vba-lineage",
                "sheets": [
                    {"name": "設定", "emit": "full", "lineage": []},
                    {"name": "グラフデータ", "emit": "full", "lineage": ["設定"]},
                    {"name": "グラフ", "emit": "full", "lineage": ["設定", "グラフデータ"]},
                ],
                "procedures": [],
            }

            markdown = Path(
                run_openpyxl(
                    str(source),
                    str(output),
                    manifest_json=json.dumps(manifest, ensure_ascii=False),
                )
            ).read_text(encoding="utf-8")

        self.assertIn("## シート: グラフ", markdown)
        self.assertIn("## シート: グラフデータ-part1", markdown)
        self.assertIn("## シート: グラフデータ-part3", markdown)
        self.assertIn("## シート: 設定", markdown)
        self.assertIn("上流シート: `設定`, `グラフデータ`", markdown)
        self.assertIn("グラフ 1", markdown)
        self.assertIn("グラフデータ シートのセル B2", markdown)

    def test_worksheet_uses_minimal_html_and_preserves_merged_cells(self) -> None:
        formulas = Workbook()
        formulas.active.merge_cells("A1:C2")
        formulas.active["A1"] = "Merged <heading>"
        formulas.active["A3"] = "line one\nline two"
        cached = Workbook()
        cached.active.merge_cells("A1:C2")
        cached.active["A1"] = "Merged <heading>"
        cached.active["A3"] = "line one\nline two"

        html, has_formulas = _render_worksheet(formulas.active, cached.active)

        self.assertFalse(has_formulas)
        self.assertIn('<td colspan="3" rowspan="2">Merged &lt;heading&gt;</td>', html)
        self.assertIn("line one<br/>line two", html)
        self.assertNotIn("style=", html)
        self.assertNotIn(" id=", html)
        self.assertNotIn("<colgroup", html)
        self.assertNotIn("| ---", html)
        formulas.close()
        cached.close()

    def test_sparse_worksheet_uses_compact_html_cell_list(self) -> None:
        formulas = Workbook()
        formulas.active["A1"] = "near"
        formulas.active["Z100"] = "far"
        cached = Workbook()
        cached.active["A1"] = "near"
        cached.active["Z100"] = "far"

        with patch.dict("os.environ", {"XLSX_MAX_TABLE_CELLS": "100"}):
            html, _ = _render_worksheet(formulas.active, cached.active)

        self.assertIn("_矩形範囲が非常に大きいため", html)
        self.assertIn("<tr><td>Z100</td><td>far</td></tr>", html)
        self.assertNotIn("| Cell | Content |", html)
        formulas.close()
        cached.close()

    def test_worksheet_is_split_at_100_rows_and_100k_characters(self) -> None:
        formulas = Workbook()
        cached = Workbook()
        for row in range(1, 102):
            formulas.active.cell(row, 1, f"row {row}")
            cached.active.cell(row, 1, f"row {row}")

        parts, _ = _render_worksheet_parts(formulas.active, cached.active, [])
        self.assertEqual([bounds[:2] for _, bounds in parts], [(1, 100), (101, 101)])
        formulas.close()
        cached.close()

        formulas = Workbook()
        cached = Workbook()
        for row in range(1, 5):
            formulas.active.cell(row, 1, "x" * 30_000)
            cached.active.cell(row, 1, "x" * 30_000)

        parts, _ = _render_worksheet_parts(formulas.active, cached.active, [])
        self.assertEqual(len(parts), 2)
        self.assertTrue(all(len(table) <= 100_000 for table, _ in parts))
        formulas.close()
        cached.close()

    def test_openpyxl_extracts_embedded_images(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            output = root / "output"

            workbook = Workbook()
            worksheet = workbook.active
            worksheet["A1"] = "Image below"
            image_stream = io.BytesIO()
            Image.new("RGB", (2, 2), "red").save(image_stream, format="PNG")
            image_stream.seek(0)
            worksheet.add_image(SpreadsheetImage(image_stream), "C3")
            workbook.save(source)
            workbook.close()

            markdown_path = Path(run_openpyxl(str(source), str(output)))
            markdown = markdown_path.read_text(encoding="utf-8")

            self.assertIn("<table>", markdown)
            self.assertNotIn("style=", markdown)
            self.assertIn("**画像位置:** Sheet シート、セル C3", markdown)
            self.assertIn(
                "![Sheet シートのセル C3 を覆う画像](media/image-1.png)",
                markdown,
            )
            self.assertTrue((output / "media" / "image-1.png").is_file())

    def test_embedded_image_reports_the_cells_covered_by_its_size(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsx"
            output = root / "output"

            workbook = Workbook()
            workbook.active["A1"] = "context"
            image_stream = io.BytesIO()
            Image.new("RGB", (240, 160), "green").save(
                image_stream,
                format="PNG",
            )
            image_stream.seek(0)
            workbook.active.add_image(SpreadsheetImage(image_stream), "C3")
            workbook.save(source)
            workbook.close()

            markdown = Path(run_openpyxl(str(source), str(output))).read_text(
                encoding="utf-8"
            )

            self.assertIn("Sheet シートのセル C3:E10 を覆う画像", markdown)
            self.assertIn("**画像位置:** Sheet シート、セル C3:E10", markdown)

    def test_openpyxl_extracts_vector_media_dropped_by_image_loader(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            initial = root / "initial.xlsx"
            source = root / "source.xlsx"
            output = root / "output"

            workbook = Workbook()
            image_stream = io.BytesIO()
            Image.new("RGB", (2, 2), "blue").save(image_stream, format="PNG")
            image_stream.seek(0)
            workbook.active.add_image(SpreadsheetImage(image_stream), "B2")
            workbook.save(initial)
            workbook.close()

            with (
                zipfile.ZipFile(initial) as original,
                zipfile.ZipFile(source, "w") as changed,
            ):
                for member in original.infolist():
                    payload = original.read(member.filename)
                    name = member.filename
                    if name == "xl/media/image1.png":
                        name = "xl/media/image1.wmf"
                        payload = b"fake-wmf-payload"
                    elif name.endswith(".xml") or name.endswith(".rels"):
                        payload = payload.replace(b"image1.png", b"image1.wmf")
                    changed.writestr(name, payload)

            with warnings.catch_warnings():
                warnings.simplefilter("ignore")
                markdown_path = Path(run_openpyxl(str(source), str(output)))
            markdown = markdown_path.read_text(encoding="utf-8")

            self.assertIn(
                "![Sheet シートのセル B2 を覆うベクター画像]"
                "(media/vector-1.wmf)",
                markdown,
            )
            self.assertIn("**画像位置:** Sheet シート、セル B2", markdown)
            self.assertEqual(
                (output / "media" / "vector-1.wmf").read_bytes(),
                b"fake-wmf-payload",
            )

    def test_libreoffice_recalculation_uses_isolated_profile(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "source.xlsm"
            output = root / "recalculated"
            source.write_bytes(make_xlsx())
            output.mkdir()
            recalculated = output / "source.xlsx"
            recalculated.write_bytes(source.read_bytes())

            completed = subprocess.CompletedProcess([], 0, "converted", "")
            with patch("formats.xlsx.subprocess.run", return_value=completed) as invoke:
                result = recalculate_with_libreoffice(
                    str(source),
                    str(output),
                    ["libreoffice"],
                )

        self.assertEqual(result, str(recalculated))
        args, options = invoke.call_args
        self.assertIn("--headless", args[0])
        self.assertIn("--convert-to", args[0])
        self.assertTrue(
            any(value.startswith("-env:UserInstallation=file:") for value in args[0])
        )
        self.assertEqual(options["timeout"], 300.0)

    async def test_recalculation_parse_images_and_llm_use_correct_resources(self) -> None:
        workers = FakeWorkers()
        options = ParseOptions(
            llm_base_url="http://override/v1",
            llm_api_key="override-key",
            llm_model="override-model",
        )
        FakeLLMClient.configurations.clear()

        with (
            patch("formats.xlsx.find_libreoffice_command", return_value=["libreoffice"]),
            patch("formats.xlsx.LLMClient", FakeLLMClient),
        ):
            result = await XlsxParser().parse(make_xlsx(), options, workers)

        self.assertEqual(
            workers.external_functions,
            [recalculate_with_libreoffice, run_openpyxl],
        )
        self.assertEqual(workers.network_calls, 1)
        self.assertEqual(result.parser, "xlsx")
        self.assertEqual(result.image_count, 1)
        self.assertIn("data:image/png;base64,", result.markdown)
        self.assertIn(
            "Description of Summary シートのセル D4 を覆う画像",
            result.markdown,
        )

    async def test_manifested_xlsm_never_opens_libreoffice(self) -> None:
        workers = FakeWorkers()
        options = ParseOptions(
            filename="source.xlsm",
            describe_images=False,
            manifest={"mode": "xlsm-vba-lineage", "sheets": [{"name": "Summary", "emit": "full"}]},
        )
        with patch("formats.xlsx.find_libreoffice_command", return_value=["libreoffice"]):
            await XlsxParser().parse(make_xlsx(), options, workers)

        self.assertEqual(workers.external_functions, [run_openpyxl])


if __name__ == "__main__":
    unittest.main()
