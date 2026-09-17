from __future__ import annotations

import io
import shutil
import tempfile
import unittest
from pathlib import Path
from typing import ClassVar
from unittest.mock import patch

from pptx import Presentation
from pptx.util import Inches
from PIL import Image

from client.llm import SlideDescriptionReview
from formats import detect
from formats.base import ParseOptions, ParseProfile
from formats.pptx import (
    PptxParser,
    _should_describe_picture,
    render_slides_with_libreoffice,
    run_pptx,
)

_PNG = io.BytesIO()
Image.new("RGB", (4, 4), "blue").save(_PNG, format="PNG")
_PNG_BYTES = _PNG.getvalue()


def make_pptx() -> bytes:
    prs = Presentation()
    title_layout, blank_layout = prs.slide_layouts[0], prs.slide_layouts[6]

    s1 = prs.slides.add_slide(title_layout)
    s1.shapes.title.text = "Quarterly Review"
    s1.placeholders[1].text = "FY2024"

    s2 = prs.slides.add_slide(blank_layout)
    box = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(3))
    tf = box.text_frame
    tf.text = "Headline point"
    p = tf.add_paragraph()
    p.text = "sub point"
    p.level = 1

    table = s2.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(4), Inches(1)).table
    table.cell(0, 0).text = "Metric"
    table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Revenue"
    table.cell(1, 1).text = "42"

    s2.shapes.add_picture(io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(1), Inches(1))
    s2.notes_slide.notes_text_frame.text = "Remember to mention hiring"

    stream = io.BytesIO()
    prs.save(stream)
    return stream.getvalue()


class FakeWorkers:
    def __init__(self) -> None:
        self.network_calls = 0
        self.external_functions = []

    async def run_external(self, fn, *args):
        self.external_functions.append(fn)
        if fn is render_slides_with_libreoffice:
            media = Path(args[1]) / "media"
            media.mkdir(parents=True, exist_ok=True)
            for index in (1, 2):
                (media / f"slide-{index}.png").write_bytes(_PNG_BYTES)
            return ["slide-1.png", "slide-2.png"]
        return fn(*args)

    async def run_network(self, fn, *args, **kwargs):
        self.network_calls += 1
        return await fn(*args, **kwargs)


class FakeLLMClient:
    configurations: ClassVar[list[dict]] = []
    calls: ClassVar[list[tuple[str, str]]] = []
    feedbacks: ClassVar[list[str]] = []
    slide_attempts: ClassVar[dict[str, int]] = {}

    def __init__(self, **kwargs) -> None:
        self.configurations.append(kwargs)

    async def describe_image(self, data_url: str, alt_text: str) -> str:
        self.calls.append(("image", alt_text))
        return f"desc[{alt_text}]"

    async def describe_slide(
        self,
        data_url: str,
        context: str,
        revision_history: list[tuple[str, str]] | None = None,
    ) -> str:
        self.calls.append(("slide", context))
        feedback = revision_history[-1][1] if revision_history else ""
        self.feedbacks.append(feedback)
        attempt = self.slide_attempts.get(context, 0) + 1
        self.slide_attempts[context] = attempt
        return f"synthesis-{attempt}[{context}]"

    async def judge_slide_description(
        self,
        data_url: str,
        context: str,
        candidate: str,
    ) -> SlideDescriptionReview:
        self.calls.append(("judge", candidate))
        attempt = int(candidate.split("-", 1)[1].split("[", 1)[0])
        scores = {1: 30.0, 2: 95.0, 3: 70.0, 4: 80.0}
        return SlideDescriptionReview(scores[attempt], f"feedback-{attempt}")

    async def close(self) -> None:
        return None


class PptxParserTests(unittest.IsolatedAsyncioTestCase):
    def test_detects_pptx_without_claiming_docx_or_xlsx(self) -> None:
        data = make_pptx()
        self.assertTrue(PptxParser.detect(data))
        self.assertIs(detect(data), PptxParser)
        self.assertFalse(PptxParser.detect(b"PK\x03\x04 not a pptx"))

    def test_run_pptx_renders_slides_titles_tables_and_extracts_pictures(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            src = root / "d.pptx"
            src.write_bytes(make_pptx())
            md_path = run_pptx(
                str(src),
                str(root / "out"),
                ["slide-1.png", "slide-2.png"],
                ParseProfile.LLM_WIKI,
            ).markdown_path
            md = md_path.read_text(encoding="utf-8")

            self.assertIn("## スライド 1", md)
            self.assertIn("### Quarterly Review", md)
            self.assertIn("## スライド 2", md)
            self.assertIn("- sub point", md)
            self.assertIn("| Metric | Value |", md)
            self.assertIn("| Revenue | 42 |", md)
            self.assertRegex(md, r"!\[[^\]]*\]\(media/image-1\.png\)")
            self.assertIn(
                "**画像位置:** 左 10.0%、上 13.3%、幅 10.0%、高さ 13.3%",
                md,
            )
            self.assertIn("> **スピーカーノート:**", md)
            self.assertIn("Remember to mention hiring", md)
            self.assertIn("### スライド全体", md)
            self.assertIn(
                "![スライド 2 全体のレンダリング](media/slide-2.png)\n\n---",
                md,
            )
            self.assertTrue((md_path.parent / "media" / "image-1.png").is_file())

    def test_run_pptx_reuses_identical_picture_payloads(self) -> None:
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(
            io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(1), Inches(1)
        )
        slide.shapes.add_picture(
            io.BytesIO(_PNG_BYTES), Inches(3), Inches(1), Inches(1), Inches(1)
        )
        stream = io.BytesIO()
        prs.save(stream)

        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "duplicate.pptx"
            source.write_bytes(stream.getvalue())
            markdown_path = run_pptx(
                str(source), str(root / "out"), profile=ParseProfile.LLM_WIKI
            ).markdown_path
            markdown = markdown_path.read_text(encoding="utf-8")

            self.assertEqual(markdown.count("(media/image-1.png)"), 2)
            self.assertEqual(
                list((markdown_path.parent / "media").iterdir()),
                [markdown_path.parent / "media" / "image-1.png"],
            )

    def test_tiny_picture_fragments_skip_individual_description(self) -> None:
        self.assertFalse(
            _should_describe_picture("Icon（左 1.0%、上 1.0%、幅 4.0%、高さ 4.0%）")
        )
        self.assertTrue(
            _should_describe_picture("Chart（左 1.0%、上 1.0%、幅 20.0%、高さ 20.0%）")
        )

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.shapes.add_picture(
            io.BytesIO(_PNG_BYTES), Inches(1), Inches(1), Inches(0.2), Inches(0.2)
        )
        stream = io.BytesIO()
        prs.save(stream)
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "tiny.pptx"
            source.write_bytes(stream.getvalue())
            markdown_path = run_pptx(
                str(source), str(root / "out"), profile=ParseProfile.LLM_WIKI
            ).markdown_path
            markdown = markdown_path.read_text(encoding="utf-8")

            self.assertNotIn("**画像位置:**", markdown)
            self.assertNotIn("media/image-", markdown)

    @unittest.skipUnless(
        shutil.which("libreoffice") or shutil.which("soffice"),
        "LibreOffice is not installed",
    )
    def test_real_libreoffice_renders_each_slide_to_png(self) -> None:
        executable = shutil.which("libreoffice") or shutil.which("soffice")
        assert executable is not None
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            source = root / "slides.pptx"
            output = root / "output"
            source.write_bytes(make_pptx())

            filenames = render_slides_with_libreoffice(
                str(source),
                str(output),
                [executable],
            )

            self.assertEqual(filenames, ["slide-1.png", "slide-2.png"])
            for filename in filenames:
                with Image.open(output / "media" / filename) as rendered:
                    self.assertEqual(rendered.format, "PNG")
                    self.assertGreaterEqual(rendered.width, 1500)

    async def test_pictures_flow_through_description_pipeline(self) -> None:
        workers = FakeWorkers()
        FakeLLMClient.configurations.clear()
        FakeLLMClient.calls.clear()
        FakeLLMClient.feedbacks.clear()
        FakeLLMClient.slide_attempts.clear()
        with (
            patch("formats.pptx.find_libreoffice_command", return_value=["libreoffice"]),
            patch("formats.pptx.LLMClient", FakeLLMClient),
        ):
            result = await PptxParser().parse(
                make_pptx(),
                ParseOptions(images=True, describe_images=True, profile=ParseProfile.LLM_WIKI),
                workers,
            )
        self.assertEqual(result.parser, "pptx")
        self.assertEqual(
            workers.external_functions,
            [render_slides_with_libreoffice, run_pptx],
        )
        self.assertEqual(result.image_count, 3)
        self.assertEqual(workers.network_calls, 3)
        self.assertIn("data:image/png;base64,", result.markdown)
        self.assertIn("desc[Picture 3（左 10.0%、上 13.3%", result.markdown)
        self.assertIn("synthesis-2[## スライド 2", result.markdown)
        slide_contexts = [
            value for kind, value in FakeLLMClient.calls if kind == "slide"
        ]
        self.assertTrue(any("Headline point" in context for context in slide_contexts))
        self.assertTrue(any("desc[Picture 3" in context for context in slide_contexts))
        self.assertTrue(all("data:image" not in context for context in slide_contexts))
        self.assertEqual(
            [kind for kind, _ in FakeLLMClient.calls].count("slide"),
            5,
        )
        self.assertEqual(
            [kind for kind, _ in FakeLLMClient.calls].count("judge"),
            4,
        )
        self.assertIn("feedback-1", FakeLLMClient.feedbacks)


if __name__ == "__main__":
    unittest.main()
