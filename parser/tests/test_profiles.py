from __future__ import annotations

import base64
import io
import json
import tempfile
import unittest
from pathlib import Path
from typing import ClassVar
from unittest.mock import patch

from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SpreadsheetImage

import server
from formats.base import ExtractedDocument, ParseOptions, ParseProfile
from formats.xlsx import RenderedWorkbook, XlsxParser, run_openpyxl
from utils.markdown_images import embed_markdown_data_urls

_PNG_BYTES = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mNk+A8AAQUBAScY42YAAAAASUVORK5CYII="
)


def _make_workbook_with_image() -> bytes:
    stream = io.BytesIO()
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Main"
    sheet["A1"] = "value"
    image_stream = io.BytesIO()
    from PIL import Image

    Image.new("RGB", (2, 2), "red").save(image_stream, format="PNG")
    image_stream.seek(0)
    sheet.add_image(SpreadsheetImage(image_stream), "B2")
    second = workbook.create_sheet("Notes")
    second["A1"] = "second-sheet"
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


class _NoLLM:
    calls: ClassVar[list[dict]] = []

    def __init__(self, **kwargs) -> None:
        _NoLLM.calls.append(kwargs)
        raise AssertionError("generic profile must never create an LLM client")

    async def close(self) -> None:  # pragma: no cover - never invoked
        return None


class GenericXlsxProfileTests(unittest.IsolatedAsyncioTestCase):
    async def test_generic_xlsm_skips_libreoffice_recalculation(self) -> None:
        functions = []

        async def run_external(fn, *args):
            functions.append(fn)
            return fn(*args)

        workers = type("Workers", (), {"run_external": staticmethod(run_external)})()
        with (
            patch("formats.xlsx.has_vba", return_value=True),
            patch("formats.xlsx.find_libreoffice_command", return_value=["libreoffice"]),
        ):
            await XlsxParser().parse(
                _make_workbook_with_image(),
                ParseOptions(
                    filename="source.xlsm",
                    profile=ParseProfile.GENERIC,
                ),
                workers,
            )

        self.assertEqual(functions, [run_openpyxl])

    async def test_generic_xlsx_keeps_every_sheet_whole_and_never_calls_llm(self) -> None:
        data = _make_workbook_with_image()

        async def run_external(fn, *args):
            return fn(*args)

        workers = type("Workers", (), {"run_external": staticmethod(run_external)})()
        with patch("formats.xlsx.LLMClient", _NoLLM):
            result = await XlsxParser().parse(
                data,
                ParseOptions(profile=ParseProfile.GENERIC),
                workers,
            )

        # No LLM client ever constructed on the generic path.
        self.assertEqual(_NoLLM.calls, [])
        # Pages: one per worksheet, unsplit, no VBA markers.
        self.assertEqual(len(result.pages), 2)
        for page in result.pages:
            self.assertNotIn("-part", page)
            self.assertNotIn("<!-- sheet-context:", page)
            self.assertNotIn("vba://", page)
        # Generic Markdown embeds images as data URLs, never as image-unit tags.
        self.assertNotIn("<image-unit>", result.markdown)
        self.assertNotIn("<image-description>", result.markdown)
        self.assertIn("](data:image/png;base64,", result.markdown)
        self.assertEqual(result.image_count, 1)

    async def test_generic_images_false_keeps_alt_text(self) -> None:
        data = _make_workbook_with_image()

        async def run_external(fn, *args):
            return fn(*args)

        workers = type("Workers", (), {"run_external": staticmethod(run_external)})()
        with patch("formats.xlsx.LLMClient", side_effect=AssertionError):
            result = await XlsxParser().parse(
                data,
                ParseOptions(profile=ParseProfile.GENERIC, images=False),
                workers,
            )

        self.assertNotIn("data:", result.markdown)
        self.assertNotIn("<image-unit>", result.markdown)
        # Alt / worksheet-location context survives as readable text.
        self.assertIn("Main シートのセル B2 を覆う画像", result.markdown)
        self.assertEqual(result.image_count, 0)
        for page in result.pages:
            self.assertNotIn("<image-media>", page)
            self.assertNotIn("data:", page)


class RunOpenpyxlReturnTests(unittest.TestCase):
    def test_run_openpyxl_returns_rendered_workbook(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            path = root / "small.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "hello"
            workbook.save(path)
            workbook.close()

            rendered = run_openpyxl(str(path), str(root / "out"))
            self.assertIsInstance(rendered, RenderedWorkbook)
            self.assertTrue(rendered.markdown.startswith("#"))


class ExtractedDocumentShapeTests(unittest.TestCase):
    def test_defaults(self) -> None:
        doc = ExtractedDocument(markdown="body")
        self.assertEqual(doc.markdown, "body")
        self.assertEqual(doc.pages, [])
        self.assertIsNone(doc.markdown_path)
        self.assertIsNone(doc.asset_root)

    def test_generic_embedding_handles_pandoc_html_images(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            (root / "media").mkdir()
            (root / "media" / "image.png").write_bytes(_PNG_BYTES)
            markdown_path = root / "document.md"

            result = embed_markdown_data_urls(
                '<img src="media/image.png" alt="diagram" width="30">',
                markdown_path,
                root,
            )

        self.assertTrue(result.startswith("![diagram](data:image/png;base64,"))
        self.assertNotIn("media/image.png", result)


# --- Server route tests -----------------------------------------------------


def _csv_body() -> bytes:
    return b"a,b\n1,2\n"


def _pptx_bytes() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    title_layout = prs.slide_layouts[0]
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(title_layout)
    s1.shapes.title.text = "Title Slide"
    s1.placeholders[1].text = "Subtitle text"
    s2 = prs.slides.add_slide(blank)
    box = s2.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1))
    box.text_frame.text = "Headline point"
    s2.shapes.add_picture(
        io.BytesIO(_PNG_BYTES), Inches(1), Inches(2), Inches(2), Inches(2)
    )
    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


class GenericPptxTests(unittest.IsolatedAsyncioTestCase):
    """Generic PPTX must emit slide text plus exactly one complete-slide PNG."""

    async def test_generic_pptx_emits_text_plus_one_screenshot_per_slide(
        self,
    ) -> None:
        import shutil

        from formats.pptx import PptxParser

        if not (shutil.which("libreoffice") or shutil.which("soffice")):
            self.skipTest("LibreOffice is required for generic PPTX screenshotting")

        async def run_external(fn, *args, **kwargs):
            return fn(*args, **kwargs)

        workers = type("Workers", (), {"run_external": staticmethod(run_external)})()

        with patch("formats.pptx.LLMClient", side_effect=AssertionError):
            result = await PptxParser().parse(
                _pptx_bytes(),
                ParseOptions(profile=ParseProfile.GENERIC),
                workers,
            )

        self.assertEqual(result.parser, "pptx")
        self.assertEqual(len(result.pages), 2)
        for slide_page in result.pages:
            self.assertRegex(slide_page, r"!\[スライド[^\]]+\]\(data:image/png;base64,")
            # exactly one image per slide page (the complete-slide screenshot)
            self.assertEqual(slide_page.count("![スライド"), 1)
        # Individual picture extraction (and the tiny-icon size filter that
        # accompanies it) must never run in generic mode.
        self.assertNotIn("**画像位置:**", result.markdown)
        self.assertNotIn("media/image-", result.markdown)
        for slide_page in result.pages:
            self.assertNotIn("media/image-", slide_page)
        # llm-wiki marker types must never appear in generic output.
        self.assertNotIn("<image-unit>", result.markdown)
        self.assertNotIn("<image-description>", result.markdown)
        self.assertIn("Headline point", result.pages[1])
        self.assertIn("Subtitle text", result.pages[0])

    async def test_generic_pptx_requires_rendering(self) -> None:
        # Generic PPTX always requires one screenshot per slide: if slide
        # rendering is unavailable the parse must fail rather than emit a
        # screenshot-less deck.
        from formats.pptx import PptxParser, SlideRenderError

        async def external(fn, *args):
            return fn(*args)

        workers = type("Workers", (), {"run_external": staticmethod(external)})()
        with (
            patch("formats.pptx.find_libreoffice_command", return_value=None),
            patch("formats.pptx.LLMClient", side_effect=AssertionError),
        ):
            with self.assertRaises(SlideRenderError):
                await PptxParser().parse(
                    _pptx_bytes(),
                    ParseOptions(profile=ParseProfile.GENERIC),
                    workers,
                )


class _FakeWorkers:
    async def run_external(self, fn, *args, **kwargs):
        return fn(*args, **kwargs)

    async def run_gpu(self, fn, *args, **kwargs):  # pragma: no cover
        return fn(*args, **kwargs)

    async def run_network(self, fn, *args, **kwargs):  # pragma: no cover
        return await fn(*args, **kwargs)

    def stats(self):  # pragma: no cover
        return {}


class ProfileRoutesTests(unittest.TestCase):
    def setUp(self) -> None:
        patcher = patch("server.URL_PREFIX", "")
        patcher.start()
        self.addCleanup(patcher.stop)
        server.app.state.workers = _FakeWorkers()
        server.app.state.mineru_api = None
        self.client = TestClient(server.app)

    def test_generic_route_rejects_manifest(self) -> None:
        response = self.client.post(
            "/parse",
            files={"file": ("x.csv", _csv_body(), "text/csv")},
            data={"manifest": json.dumps({"mode": "anything"})},
        )
        self.assertEqual(response.status_code, 400)
        self.assertIn("manifest is only supported", response.json()["detail"])

    def test_generic_route_returns_empty_pages_for_csv(self) -> None:
        response = self.client.post(
            "/parse",
            files={"file": ("x.csv", _csv_body(), "text/csv")},
        )
        self.assertEqual(response.status_code, 200)
        payload = response.json()
        self.assertEqual(payload["parser"], "csv")
        self.assertEqual(payload["pages"], [])
        self.assertIsInstance(payload["markdown"], str)

    def test_llm_wiki_route_accepts_headers_and_still_rejects_unsupported(self) -> None:
        response = self.client.post(
            "/parse/llm-wiki",
            files={"file": ("x.txt", b"hello world", "text/plain")},
            headers={"X-LLM-Base-URL": "http://llm/v1", "X-LLM-Model": "m"},
        )
        self.assertEqual(response.status_code, 415)

    def test_generic_route_ignores_describe_images(self) -> None:
        # A generic parse with describe_images=true must not attempt an LLM call
        # and must not error out. Reuse CSV since there is no image path.
        response = self.client.post(
            "/parse?describe_images=true",
            files={"file": ("x.csv", _csv_body(), "text/csv")},
        )
        self.assertEqual(response.status_code, 200)


if __name__ == "__main__":
    unittest.main()
