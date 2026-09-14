"""Regenerate xlsx_kitchen_sink.xlsx.

A single workbook that touches every branch in formats/xlsx.py:
cross-sheet formulas with no stored cached value, a dense table, a sparse
sheet that forces the cell-list view, two embedded PNG images at different
anchors, an empty sheet, and pipe/newline/unicode escaping.
"""

from __future__ import annotations

import io
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from PIL import Image, ImageDraw


def _png(color: str, label: str) -> io.BytesIO:
    im = Image.new("RGB", (240, 160), color)
    draw = ImageDraw.Draw(im)
    draw.rectangle([8, 8, 231, 151], outline="white", width=3)
    draw.text((20, 70), label, fill="white")
    buf = io.BytesIO()
    im.save(buf, format="PNG")
    buf.seek(0)
    return buf


def build() -> Workbook:
    wb = Workbook()

    summary = wb.active
    summary.title = "Summary"
    summary["A1"] = "Metric"
    summary["B1"] = "Value"
    summary["A2"] = "Report date"
    summary["B2"] = datetime(2024, 6, 30, 15, 30)
    summary["A3"] = "Quarter start"
    summary["B3"] = date(2024, 4, 1)
    summary["A4"] = "Total revenue (Data!)"
    summary["B4"] = "=SUM(Data!B2:B13)"
    summary["A5"] = "Average revenue"
    summary["B5"] = "=AVERAGE(Data!B2:B13)"
    summary["A6"] = "Max month"
    summary["B6"] = "=MAX(Data!B2:B13)"
    summary["A7"] = "Pipe | and newline"
    summary["B7"] = "line one\nline two"
    summary["A8"] = "Unicode"
    summary["B8"] = "Ω ≈ 3.14 — café —"
    summary["A9"] = "Nested formula"
    summary["B9"] = '=IF(B4>1000,"big","small")'

    data = wb.create_sheet("Data")
    data.append(["Month", "Revenue", "Cost", "Margin"])
    for month in range(1, 13):
        row = month + 1
        data.append([f"2024-{month:02d}", 100 * month + 50, None, None])
        data[f"C{row}"] = f"=B{row}*0.6"
        data[f"D{row}"] = f"=B{row}-C{row}"

    sparse = wb.create_sheet("Sparse")
    sparse["A1"] = "top-left anchor"
    sparse["B2"] = 42
    sparse["ZZ4000"] = "bottom-right far cell"  # area ~2.8M cells >> XLSX_MAX_TABLE_CELLS
    sparse["AA10"] = "=B2*10"

    pictures = wb.create_sheet("Pictures")
    pictures["A1"] = "Two embedded images below"
    pictures.add_image(XLImage(_png("navy", "CHART ALPHA")), "B3")
    pictures.add_image(XLImage(_png("darkgreen", "DIAGRAM BETA")), "F14")

    wb.create_sheet("EmptySheet")
    return wb


if __name__ == "__main__":
    target = Path(__file__).with_name("xlsx_kitchen_sink.xlsx")
    build().save(target)
    print(f"wrote {target}")
